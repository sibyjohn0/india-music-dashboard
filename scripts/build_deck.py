#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_deck.py — the Indie Music India artist working-template kit (.pptx).

A consultant-grade, on-brand (poppy neo-brutalist) Google Slides kit covering the
whole IIM engagement pipeline: Identity, Positioning, Social, Release, Promotion,
Distribution — plus a drag-and-drop Component Library and reusable variations.

Build .pptx locally with full brand control (colours, Bricolage/Inter/Space Mono,
hard offset shadows, media placeholders), then upload to Drive converted to native
Google Slides so a teammate can duplicate a slide, swap text/media, and delete
what they don't need.

Output: social/decks/IIM-Artist-Working-Deck.pptx  (gitignored)
Usage:  python3 scripts/build_deck.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

REPO = Path(__file__).resolve().parent.parent
MARK = REPO / "assets" / "brand" / "mark.png"
OUT = REPO / "social" / "decks" / "IIM-Artist-Working-Deck.pptx"

# ---- brand tokens -----------------------------------------------------------
def C(h): return RGBColor.from_string(h)
CREAM="FFF7EE"; INK="241B2E"; PINK="FF4D8D"; YELLOW="FFD23F"; MINT="1FCF9E"
VIOLET="8B5CF6"; WHITE="FFFFFF"; MUT="6B6076"; NEUT="F1ECF5"; CARD="FFFFFF"
HEAD="Bricolage Grotesque"; BODY="Inter"; MONO="Space Mono"
SW, SH = 13.333, 7.5
MX = 0.7

SECTIONS = [  # (num, title, accent, blurb)
 ("01","IDENTITY",PINK,"Who the artist actually is."),
 ("02","POSITIONING",VIOLET,"Where they sit in a crowded market."),
 ("03","SOCIAL",YELLOW,"The always-on content engine."),
 ("04","RELEASE",MINT,"Getting the music out, on schedule."),
 ("05","PROMOTION",PINK,"Getting it heard by the right people."),
 ("06","DISTRIBUTION",VIOLET,"Where it lives, and where the money is."),
]

prs = Presentation()
prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

# ---- primitives -------------------------------------------------------------
def _noshadow(shp): shp.shadow.inherit = False

def slide(bg=CREAM, bar=PINK):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = C(bg)
    if bar:
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(0.14))
        b.fill.solid(); b.fill.fore_color.rgb = C(bar); b.line.fill.background(); _noshadow(b)
    return s

def _round(shp, radius=0.09):
    try: shp.adjustments[0] = radius
    except Exception: pass

def card(s, x, y, w, h, fill=CARD, line=INK, lw=2.0, radius=0.07, shadow=INK, off=0.07):
    if shadow:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+off), Inches(y+off), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = C(shadow); sh.line.fill.background(); _round(sh, radius); _noshadow(sh)
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = C(fill)
    if line: c.line.color.rgb = C(line); c.line.width = Pt(lw)
    else: c.line.fill.background()
    _round(c, radius); _noshadow(c)
    return c

def _set(run, font, size, color, bold, italic=False, spacing=None):
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
    run.font.italic = italic; run.font.color.rgb = C(color)
    if spacing is not None:
        rPr = run._r.get_or_add_rPr(); rPr.set('spc', str(int(spacing*100)))

def text(s, x, y, w, h, body, font=BODY, size=16, color=INK, bold=False, italic=False,
         align="l", anchor="t", leading=1.06, sp_after=4, spacing=None, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.vertical_anchor = {"t":MSO_ANCHOR.TOP,"m":MSO_ANCHOR.MIDDLE,"b":MSO_ANCHOR.BOTTOM}[anchor]
    al = {"l":PP_ALIGN.LEFT,"c":PP_ALIGN.CENTER,"r":PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al; p.line_spacing = leading; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        r = p.add_run(); r.text = line; _set(r, font, size, color, bold, italic, spacing)
    return tb

def chip(s, x, y, label, fill=YELLOW, fg=INK, size=11):
    w = max(0.5, len(label)*size*0.0095 + 0.36)  # tuned to Space Mono advance
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    c.fill.solid(); c.fill.fore_color.rgb = C(fill); c.line.color.rgb = C(INK); c.line.width = Pt(1.5)
    _round(c, 0.5); _noshadow(c)
    tf = c.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; _set(r, MONO, size, fg, True, spacing=0.5)
    return w

def bullets(s, x, y, w, h, items, marker="→", size=15, color=INK, gap=7, mfill=None, font=BODY, leading=1.06):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = leading; p.space_after = Pt(gap); p.space_before = Pt(0)
        rm = p.add_run(); rm.text = marker + "  "; _set(rm, font, size, mfill or PINK, True)
        rr = p.add_run(); rr.text = it; _set(rr, font, size, color, False)
    return tb

def media(s, x, y, w, h, label, ratio="", fill=NEUT):
    c = card(s, x, y, w, h, fill=fill, line=INK, lw=1.75, radius=0.05, shadow=None)
    # dashed inner feel via label only; centered instruction
    t = text(s, x, y+h/2-0.5, w, 1.0, ("▶  "+label) + (("\n"+ratio) if ratio else ""),
             font=MONO, size=12, color=MUT, align="c", anchor="m")
    return c

def footer(s, label):
    if MARK.exists():
        s.shapes.add_picture(str(MARK), Inches(MX), Inches(SH-0.52), height=Inches(0.28))
    text(s, MX+0.36, SH-0.55, 8.0, 0.34, f"indiemusicindia.com     ·     {label}", font=MONO, size=9, color=MUT, anchor="m")

def stamp_numbers():
    """Global pass: number every slide bottom-right so it's easy to reference.
    A yellow pill reads on both cream work slides and dark dividers."""
    for i, s in enumerate(prs.slides, 1):
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(SW-1.02), Inches(SH-0.64), Inches(0.62), Inches(0.42))
        c.fill.solid(); c.fill.fore_color.rgb = C(YELLOW); c.line.color.rgb = C(INK); c.line.width = Pt(1.75)
        _round(c, 0.5); _noshadow(c)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        tf.margin_top = 0; tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = f"{i:02d}"; _set(r, MONO, 13, INK, True)

def divider(num, title, accent, blurb, inside):
    s = slide(bg=INK, bar=accent)
    text(s, MX, 1.5, 3.0, 2.0, num, font=HEAD, size=150, color=accent, bold=True)
    text(s, MX, 3.55, 11.0, 1.4, title, font=HEAD, size=68, color=CREAM, bold=True)
    text(s, MX, 4.85, 9.0, 0.6, blurb, font=BODY, size=20, color="D9CFE6")
    # "what's inside" chips
    x = MX
    for it in inside:
        w = chip(s, x, 5.7, it, fill=accent, fg=INK, size=12); x += w + 0.18
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def head(s, kicker, title, accent=PINK, y=0.55):
    chip(s, MX, y, kicker, fill=accent)
    text(s, MX, y+0.42, 11.9, 0.9, title, font=HEAD, size=34, color=INK, bold=True)

# =============================================================================
# SECTION 0 — COVER + ORIENTATION
# =============================================================================
def s_cover():
    s = slide(bg=CREAM, bar=None)
    # big color band top
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(0.22))
    band.fill.solid(); band.fill.fore_color.rgb = C(PINK); band.line.fill.background(); _noshadow(band)
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(MX), Inches(0.85), height=Inches(0.9))
    chip(s, MX, 2.05, "ARTIST WORKING DECK  ·  INTERNAL KIT", fill=YELLOW)
    text(s, MX, 2.55, 12.0, 2.6, "The Indie Music\nIndia Playbook Deck", font=HEAD, size=66, color=INK, bold=True, leading=0.98)
    text(s, MX, 5.15, 11.0, 0.9, "One template kit to take an artist from identity to distribution.\nDuplicate a variation, swap the text and media, ship the work.",
         font=BODY, size=19, color=MUT, leading=1.15)
    x = MX
    for lbl, col in [("IDENTITY",PINK),("POSITIONING",VIOLET),("SOCIAL",YELLOW),("RELEASE",MINT),("PROMOTION",PINK),("DISTRIBUTION",VIOLET)]:
        w = chip(s, x, 6.15, lbl, fill=col, size=9); x += w + 0.12
    footer(s, "COVER")

def s_howto():
    s = slide(); head(s, "READ ME FIRST", "How to use this kit", accent=MINT)
    steps = [
        ("1  Find the section","Six colour-coded sections follow the IIM pipeline. Section dividers are dark; work slides are cream."),
        ("2  Pick a variation","Most topics ship in 2-3 layouts (text-led, media-led, worksheet). Duplicate the one that fits the artist."),
        ("3  Swap, don't rebuild","Click any block, replace the placeholder text or drop your media into the grey frames. Keep the brand styling."),
        ("4  Raid the library","The Component Library at the back holds ready blocks: stat, quote, checklist, matrix, table. Copy onto any slide."),
        ("5  Delete the rest","This is a superset. Cut every slide the artist doesn't need. A tight 12-slide deck beats a padded 40."),
        ("6  Keep it on-brand","Colours and fonts are locked to the site. If a font shows as a substitute, add Bricolage Grotesque in Slides (it's a Google Font)."),
    ]
    cw, ch, gx, gy = 3.72, 1.78, 0.28, 0.28
    x0, y0 = MX, 1.72
    cols = [PINK, VIOLET, YELLOW, MINT, PINK, VIOLET]
    for i,(t,d) in enumerate(steps):
        r, cx = divmod(i, 3)
        x = x0 + cx*(cw+gx); y = y0 + r*(ch+gy)
        card(s, x, y, cw, ch, fill=CARD)
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.14), Inches(ch))
        bar.fill.solid(); bar.fill.fore_color.rgb = C(cols[i]); bar.line.fill.background(); _round(bar,0.5); _noshadow(bar)
        text(s, x+0.3, y+0.18, cw-0.5, 0.5, t, font=HEAD, size=17, color=INK, bold=True)
        text(s, x+0.3, y+0.62, cw-0.5, ch-0.72, d, font=BODY, size=12, color=MUT, leading=1.12)
    footer(s, "ORIENTATION")

def s_brand():
    s = slide(); head(s, "BRAND QUICK-REFERENCE", "Stay on the house style", accent=VIOLET)
    # palette
    text(s, MX, 1.7, 4.0, 0.4, "PALETTE", font=MONO, size=12, color=MUT, bold=True, spacing=1)
    sw = [("Cream",CREAM,INK),("Ink",INK,CREAM),("Pink",PINK,WHITE),("Yellow",YELLOW,INK),("Mint",MINT,INK),("Violet",VIOLET,WHITE)]
    x = MX
    for name,hexv,fg in sw:
        card(s, x, 2.1, 1.28, 1.28, fill=hexv, line=INK, radius=0.08, shadow=INK, off=0.06)
        text(s, x, 2.55, 1.28, 0.4, name, font=HEAD, size=14, color=fg, bold=True, align="c")
        text(s, x, 2.95, 1.28, 0.3, "#"+hexv, font=MONO, size=8.5, color=fg, align="c")
        x += 1.45
    # type
    text(s, MX, 3.9, 6.0, 0.4, "TYPE", font=MONO, size=12, color=MUT, bold=True, spacing=1)
    text(s, MX, 4.25, 6.4, 0.7, "Bricolage Grotesque", font=HEAD, size=30, color=INK, bold=True)
    text(s, MX, 4.9, 6.4, 0.4, "Display / headlines", font=BODY, size=12, color=MUT)
    text(s, MX, 5.35, 6.4, 0.6, "Inter: body copy, the workhorse for everything you read.", font=BODY, size=16, color=INK)
    text(s, MX, 5.95, 6.4, 0.5, "SPACE MONO: labels, chips, numbers.", font=MONO, size=13, color=INK, bold=True)
    # do / dont
    card(s, 7.7, 4.1, 4.9, 2.6, fill=CARD)
    text(s, 7.95, 4.28, 4.5, 0.4, "DO  /  DON'T", font=MONO, size=12, color=MUT, bold=True, spacing=1)
    bullets(s, 7.95, 4.7, 4.4, 1.0, ["Keep 2px ink borders and hard shadows","Let one accent colour lead a slide","Use real numbers and specifics"], marker="✓", mfill=MINT, size=12.5, gap=4)
    bullets(s, 7.95, 5.75, 4.4, 1.0, ["Don't add gradients or soft drop-shadows","Don't mix three accent colours in one block","Don't centre everything or use stock clichés"], marker="✗", mfill=PINK, size=12.5, gap=4)
    footer(s, "ORIENTATION")

def s_pipeline():
    s = slide(); head(s, "THE ENGAGEMENT", "The pipeline, end to end", accent=PINK)
    text(s, MX, 1.55, 11.9, 0.5, "Every artist moves left to right. Each stage feeds the next: skip one and the ones after it get expensive.",
         font=BODY, size=15, color=MUT)
    cw, gx = 1.86, 0.14; x0, y = MX, 2.35
    for i,(num,title,col,blurb) in enumerate(SECTIONS):
        x = x0 + i*(cw+gx)
        card(s, x, y, cw, 3.2, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(0.7))
        top.fill.solid(); top.fill.fore_color.rgb = C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, y+0.1, cw, 0.5, num, font=HEAD, size=26, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x+0.14, y+0.85, cw-0.28, 0.6, title, font=HEAD, size=15, color=INK, bold=True, align="c")
        text(s, x+0.16, y+1.45, cw-0.32, 1.6, blurb, font=BODY, size=11, color=MUT, align="c", leading=1.12)
        if i < 5:
            text(s, x+cw-0.06, y+1.25, 0.4, 0.5, "›", font=HEAD, size=26, color=INK, bold=True)
    footer(s, "OVERVIEW")

def s_projectcover():
    s = slide(bg=INK, bar=YELLOW)
    chip(s, MX, 0.9, "PROJECT COVER  ·  DUPLICATE PER ARTIST", fill=YELLOW)
    text(s, MX, 1.5, 11.0, 1.6, "[ Artist Name ]", font=HEAD, size=64, color=CREAM, bold=True)
    text(s, MX, 2.95, 11.0, 0.6, "[ One-line positioning: the genre / the promise ]", font=BODY, size=22, color="D9CFE6", italic=True)
    fields = [("STRATEGIST","[ name ]"),("STAGE","[ Identity → Distribution ]"),("ENGAGEMENT","[ generic / 4-hour / monthly ]"),("START DATE","[ dd mmm yyyy ]")]
    x = MX
    for k,v in fields:
        card(s, x, 4.3, 2.85, 1.35, fill=INK, line=YELLOW, lw=1.75, shadow=None)
        text(s, x+0.22, 4.5, 2.5, 0.3, k, font=MONO, size=10, color=YELLOW, bold=True, spacing=0.5)
        text(s, x+0.22, 4.9, 2.5, 0.7, v, font=BODY, size=15, color=CREAM)
        x += 3.02
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

# =============================================================================
# SECTION 1 — IDENTITY
# =============================================================================
def s_identity_snapshot():
    s = slide(); head(s, "01 · IDENTITY  ·  VAR A (TEXT)", "Artist snapshot", accent=PINK)
    card(s, MX, 1.7, 5.7, 4.9, fill=CARD)
    rows = [("NAME","[ artist / act name ]"),("GENRE / LANE","[ e.g. Hindi indie-pop, bedroom R&B ]"),
            ("BASED IN","[ city ]"),("FORMED","[ year ]"),("FOR FANS OF","[ 3 reference artists ]"),
            ("LINKS","[ IG · Spotify · YouTube ]")]
    y = 1.95
    for k,v in rows:
        text(s, MX+0.28, y, 2.2, 0.4, k, font=MONO, size=10.5, color=MUT, bold=True, spacing=0.5)
        text(s, MX+0.28, y+0.28, 5.0, 0.5, v, font=BODY, size=15, color=INK)
        y += 0.76
    card(s, 6.75, 1.7, 5.85, 4.9, fill=NEUT, line=INK, shadow=INK)
    text(s, 7.0, 1.95, 5.4, 0.4, "THE ONE-LINER", font=MONO, size=11, color=MUT, bold=True, spacing=1)
    text(s, 7.0, 2.35, 5.4, 1.6, "“[ In one sentence, who is this artist and why should anyone care? ]”",
         font=HEAD, size=26, color=INK, bold=True, leading=1.05)
    text(s, 7.0, 4.4, 5.4, 0.4, "WHAT THEY ARE NOT", font=MONO, size=11, color=MUT, bold=True, spacing=1)
    text(s, 7.0, 4.8, 5.4, 1.4, "[ Defining by contrast sharpens identity. e.g. 'not a Bollywood playback voice, not a lo-fi loop act.' ]",
         font=BODY, size=14, color=INK, leading=1.15)
    footer(s, "IDENTITY")

def s_identity_origin():
    s = slide(); head(s, "01 · IDENTITY  ·  VAR B (NARRATIVE)", "Origin story", accent=PINK)
    text(s, MX, 1.65, 7.0, 3.9,
         "[ The narrative a journalist would actually run. Where the artist is from, the turn that "
         "made them start, the first room that got it. Two short paragraphs, no résumé bullets. ]\n\n"
         "[ Second beat: what changed, the sound they found, why now. Keep it human and specific: "
         "a place, a person, a moment, not adjectives. ]",
         font=BODY, size=15.5, color=INK, leading=1.25)
    card(s, 8.0, 1.7, 4.6, 4.5, fill=YELLOW, line=INK, shadow=INK)
    text(s, 8.25, 1.95, 4.1, 0.35, "PULL QUOTE", font=MONO, size=10.5, color=INK, bold=True, spacing=1)
    text(s, 8.25, 2.4, 4.1, 3.4, "“[ The one line from the story you'd put on a poster. ]”",
         font=HEAD, size=28, color=INK, bold=True, leading=1.05, anchor="m")
    footer(s, "IDENTITY")

def s_identity_sonic():
    s = slide(); head(s, "01 · IDENTITY  ·  DEEP-DIVE", "Sonic identity", accent=PINK)
    cols = [("SOUND",MINT,"[ instruments, production texture, tempo range, the signature move ]"),
            ("MOOD",VIOLET,"[ the feeling in the room: 3-4 adjectives a fan would actually use ]"),
            ("REFERENCES",PINK,"[ 3 artists + why. 'X's honesty, Y's groove, Z's restraint' ]")]
    cw, gx = 3.86, 0.28; x = MX
    for title,col,body in cols:
        card(s, x, 1.75, cw, 3.0, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.75), Inches(cw), Inches(0.6))
        top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, 1.83, cw, 0.45, title, font=HEAD, size=17, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x+0.24, 2.55, cw-0.48, 2.0, body, font=BODY, size=13.5, color=INK, leading=1.18)
        x += cw+gx
    media(s, MX, 5.0, 6.1, 1.5, "DROP A REFERENCE-TRACK MOODBOARD OR WAVEFORM", "16:9")
    media(s, 7.0, 5.0, 5.6, 1.5, "DROP A 15-SEC SIGNATURE-SOUND CLIP THUMBNAIL", "16:9")
    footer(s, "IDENTITY")

def s_identity_visual():
    s = slide(); head(s, "01 · IDENTITY  ·  VAR C (MEDIA)", "Visual identity board", accent=PINK)
    media(s, MX, 1.7, 6.0, 3.15, "DROP HERO PRESS SHOT", "3:2")
    media(s, 6.95, 1.7, 2.75, 1.5, "LOGO", "1:1")
    media(s, 9.85, 1.7, 2.75, 1.5, "ALT SHOT", "1:1")
    # palette strip
    text(s, 6.95, 3.35, 5.6, 0.3, "ARTIST PALETTE (SWAP TO THEIRS)", font=MONO, size=10, color=MUT, bold=True, spacing=0.5)
    x = 6.95
    for hexv in [INK,PINK,YELLOW,MINT,VIOLET]:
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.7), Inches(1.05), Inches(1.05))
        c.fill.solid(); c.fill.fore_color.rgb=C(hexv); c.line.color.rgb=C(INK); c.line.width=Pt(1.75); _round(c,0.08); _noshadow(c)
        x += 1.14
    media(s, MX, 5.05, 12.0, 1.5, "DROP TYPE SPECIMEN + GRAPHIC MOTIF STRIP", "wide")
    footer(s, "IDENTITY")

def s_identity_audit():
    s = slide(); head(s, "01 · IDENTITY  ·  WORKSHEET", "Identity audit", accent=PINK)
    text(s, MX, 1.55, 11.9, 0.4, "Score each 1-5. Anything under 3 is this month's homework.", font=BODY, size=14, color=MUT)
    items = ["A one-liner a stranger could repeat","Consistent name + handle everywhere","A press photo that looks like the music",
             "A logo / wordmark that survives a thumbnail","A clear 'for fans of' shortlist","A story with a place, a turn, a moment",
             "A defined lane, and what they're NOT","Bio in three lengths (280 / 100 / 40 words)"]
    cw, ch, gx, gy = 5.85, 0.86, 0.3, 0.24; x0, y0 = MX, 2.05
    for i,it in enumerate(items):
        r,cx = divmod(i,2); x=x0+cx*(cw+gx); y=y0+r*(ch+gy)
        card(s, x, y, cw, ch, fill=CARD, shadow=INK, off=0.05)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+0.22), Inches(y+0.24), Inches(0.38), Inches(0.38))
        box.fill.solid(); box.fill.fore_color.rgb=C(NEUT); box.line.color.rgb=C(INK); box.line.width=Pt(1.75); _round(box,0.2); _noshadow(box)
        text(s, x+0.78, y+0.13, cw-1.7, 0.6, it, font=BODY, size=12.5, color=INK, anchor="m", leading=1.05)
        text(s, x+cw-0.95, y+0.22, 0.8, 0.42, "▢ /5", font=MONO, size=12, color=MUT, align="r")
    footer(s, "IDENTITY")

# =============================================================================
# SECTION 2 — POSITIONING
# =============================================================================
def s_pos_statement():
    s = slide(); head(s, "02 · POSITIONING  ·  CORE", "The positioning statement", accent=VIOLET)
    card(s, MX, 1.9, 11.9, 2.5, fill=VIOLET, line=INK, shadow=INK)
    text(s, MX+0.4, 2.15, 11.1, 2.0,
         "For [ the specific listener ] who [ the need or moment ],\n[ artist ] is the [ category ] that [ the one benefit nobody else credibly offers ].",
         font=HEAD, size=27, color=WHITE, bold=True, leading=1.15, anchor="m")
    text(s, MX, 4.65, 11.9, 0.4, "Fill every bracket. If two artists could swap names into this line, it isn't positioning yet.", font=BODY, size=14, color=MUT)
    cols = [("AUDIENCE","[ who exactly ]"),("NEED","[ the itch ]"),("CATEGORY","[ the lane ]"),("BENEFIT","[ the wedge ]")]
    x = MX
    for k,v in cols:
        card(s, x, 5.25, 2.85, 1.15, fill=CARD)
        text(s, x+0.22, 5.42, 2.5, 0.3, k, font=MONO, size=10, color=MUT, bold=True, spacing=0.5)
        text(s, x+0.22, 5.75, 2.5, 0.6, v, font=BODY, size=14, color=INK)
        x += 3.02
    footer(s, "POSITIONING")

def s_pos_personas():
    s = slide(); head(s, "02 · POSITIONING  ·  DEEP-DIVE", "Audience personas", accent=VIOLET)
    p = [("THE SUPERFAN",PINK,["Age / city","What they already listen to","Where they hang out online","What would make them buy a ticket"]),
         ("THE DISCOVERER",MINT,["How they find new music","Which playlist / creator they trust","The hook that stops their scroll","What makes them save vs skip"]),
         ("THE GATEKEEPER",YELLOW,["Curator / journalist / booker","What lands in their inbox","The proof they need to say yes","One warm intro that reaches them"])]
    cw, gx = 3.86, 0.28; x = MX
    for name,col,rows in p:
        card(s, x, 1.75, cw, 4.7, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.75), Inches(cw), Inches(0.65))
        top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, 1.84, cw, 0.5, name, font=HEAD, size=18, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        media(s, x+0.3, 2.6, cw-0.6, 1.1, "AVATAR", "1:1")
        bullets(s, x+0.3, 3.85, cw-0.55, 2.4, rows, size=12.5, gap=8, mfill=col)
        x += cw+gx
    footer(s, "POSITIONING")

def s_pos_map():
    s = slide(); head(s, "02 · POSITIONING  ·  TOOL", "Perceptual map", accent=VIOLET)
    text(s, MX, 1.55, 11.9, 0.4, "Plot the artist and 5-6 peers. Empty quadrants are the opportunity.", font=BODY, size=14, color=MUT)
    # axes box
    ox, oy, ow, oh = MX+1.4, 2.1, 8.0, 4.3
    card(s, ox, oy, ow, oh, fill=WHITE, line=INK, shadow=None, radius=0.02)
    # cross lines
    for (x1,y1,x2,y2) in [(ox, oy+oh/2, ox+ow, oy+oh/2),(ox+ow/2, oy, ox+ow/2, oy+oh)]:
        ln = s.shapes.add_connector(2, Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        ln.line.color.rgb=C(INK); ln.line.width=Pt(1.5); _noshadow(ln)
    text(s, ox, oy-0.05, ow, 0.3, "[ TOP AXIS LABEL, e.g. Polished ]", font=MONO, size=10, color=MUT, align="c")
    text(s, ox, oy+oh-0.02, ow, 0.3, "[ BOTTOM, e.g. Raw ]", font=MONO, size=10, color=MUT, align="c")
    text(s, ox-1.35, oy+oh/2-0.15, 1.3, 0.3, "[ LEFT ]", font=MONO, size=10, color=MUT, align="r")
    text(s, ox+ow+0.1, oy+oh/2-0.15, 1.3, 0.3, "[ RIGHT ]", font=MONO, size=10, color=MUT)
    # sample dot for the artist
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ox+ow*0.62), Inches(oy+oh*0.3), Inches(0.34), Inches(0.34))
    dot.fill.solid(); dot.fill.fore_color.rgb=C(PINK); dot.line.color.rgb=C(INK); dot.line.width=Pt(2); _noshadow(dot)
    text(s, ox+ow*0.62+0.4, oy+oh*0.3-0.02, 2.0, 0.34, "YOU", font=MONO, size=11, color=INK, bold=True, anchor="m")
    card(s, 9.8, 2.1, 2.8, 4.3, fill=NEUT, line=INK, shadow=INK)
    text(s, 10.0, 2.3, 2.4, 0.35, "READOUT", font=MONO, size=10.5, color=MUT, bold=True, spacing=1)
    text(s, 10.0, 2.7, 2.45, 3.5, "[ Which quadrant is crowded?\n\nWhich is empty and true to the artist?\n\nThat empty-and-true corner is the position. ]",
         font=BODY, size=12.5, color=INK, leading=1.18)
    footer(s, "POSITIONING")

def s_pos_wedge():
    s = slide(); head(s, "02 · POSITIONING  ·  CORE", "The wedge", accent=VIOLET)
    text(s, MX, 1.8, 11.9, 1.4, "“[ The one true thing this artist has that competitors can't copy overnight. ]”",
         font=HEAD, size=40, color=INK, bold=True, leading=1.05)
    text(s, MX, 3.5, 11.9, 0.4, "THREE PROOF POINTS", font=MONO, size=12, color=MUT, bold=True, spacing=1)
    cols = [("PROOF","[ a receipt: a number, a placement, a co-sign ]",MINT),
            ("PROOF","[ a craft detail only they do ]",VIOLET),
            ("PROOF","[ a relationship or scene they own ]",PINK)]
    cw, gx = 3.86, 0.28; x = MX
    for k,v,col in cols:
        card(s, x, 4.0, cw, 2.3, fill=CARD)
        chip(s, x+0.24, 4.22, k, fill=col, size=10)
        text(s, x+0.24, 4.75, cw-0.5, 1.4, v, font=BODY, size=14.5, color=INK, leading=1.2)
        x += cw+gx
    footer(s, "POSITIONING")

def s_pos_teardown():
    s = slide(); head(s, "02 · POSITIONING  ·  WORKSHEET", "Competitor teardown", accent=VIOLET)
    cols = ["ARTIST","THEIR LANE","WHAT WORKS","THE GAP YOU FILL"]
    widths = [2.6, 3.0, 3.2, 3.1]; x0, y0 = MX, 1.85
    # header
    x = x0
    for c,w in zip(cols,widths):
        h = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y0), Inches(w), Inches(0.5))
        h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.5); _noshadow(h)
        text(s, x+0.15, y0+0.08, w-0.2, 0.35, c, font=MONO, size=11, color=CREAM, bold=True, spacing=0.5)
        x += w
    for r in range(5):
        y = y0+0.5+r*0.82; x = x0
        fill = WHITE if r%2==0 else NEUT
        for w in widths:
            cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.82))
            cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(1); _noshadow(cell)
            text(s, x+0.15, y+0.22, w-0.2, 0.5, "[ … ]", font=BODY, size=12, color=MUT)
            x += w
    footer(s, "POSITIONING")

# =============================================================================
# SECTION 3 — SOCIAL
# =============================================================================
def s_social_channels():
    s = slide(); head(s, "03 · SOCIAL  ·  STRATEGY", "Channel strategy", accent=YELLOW)
    ch = [("INSTAGRAM",PINK,"Discovery + personality. Reels to reach, stories to bond."),
          ("YOUTUBE",VIOLET,"Depth: live takes, visualisers, the searchable back-catalogue."),
          ("SPOTIFY PROFILE",MINT,"The conversion endpoint. Canvas, bio, artist pick, playlists."),
          ("WHATSAPP / DISCORD",YELLOW,"The owned list. Where superfans hear it first, no algorithm.")]
    cw, ch_h, gx, gy = 5.85, 2.1, 0.3, 0.3; x0, y0 = MX, 1.75
    for i,(name,col,body) in enumerate(ch):
        r,cx = divmod(i,2); x=x0+cx*(cw+gx); y=y0+r*(ch_h+gy)
        card(s, x, y, cw, ch_h, fill=CARD)
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.16), Inches(ch_h))
        bar.fill.solid(); bar.fill.fore_color.rgb=C(col); bar.line.fill.background(); _round(bar,0.5); _noshadow(bar)
        text(s, x+0.35, y+0.25, cw-0.6, 0.5, name, font=HEAD, size=21, color=INK, bold=True)
        text(s, x+0.35, y+0.85, cw-0.7, 1.1, body, font=BODY, size=14, color=MUT, leading=1.18)
        text(s, x+cw-1.7, y+0.28, 1.5, 0.35, "ROLE", font=MONO, size=10, color=MUT, bold=True, align="r", spacing=1)
    footer(s, "SOCIAL")

def s_social_pillars():
    s = slide(); head(s, "03 · SOCIAL  ·  ENGINE", "Content pillars", accent=YELLOW)
    text(s, MX, 1.55, 11.9, 0.4, "Four buckets on rotation. Every post belongs to one. If it doesn't, it doesn't go out.", font=BODY, size=14, color=MUT)
    p = [("THE MUSIC",PINK,"Snippets, live takes, the making-of. The reason they're here."),
         ("THE PERSON",VIOLET,"Face, voice, day-to-day. Turns a track into a relationship."),
         ("THE VALUE",MINT,"Teach one useful thing. Saves and shares, reaches new people."),
         ("THE PROOF",YELLOW,"Crowds, press, placements, fan reactions. Quiet credibility.")]
    cw, gx = 2.9, 0.14; x = MX
    for name,col,body in p:
        card(s, x, 2.2, cw, 3.9, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(cw), Inches(1.1))
        top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x+0.2, 2.45, cw-0.4, 0.7, name, font=HEAD, size=19, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, anchor="m")
        text(s, x+0.25, 3.5, cw-0.5, 2.3, body, font=BODY, size=13.5, color=INK, leading=1.2)
        text(s, x+0.25, 5.65, cw-0.5, 0.35, "≈ 1 in 4 posts", font=MONO, size=10, color=MUT, bold=True)
        x += cw+gx
    footer(s, "SOCIAL")

def s_social_cadence():
    s = slide(); head(s, "03 · SOCIAL  ·  WORKSHEET", "Weekly cadence", accent=YELLOW)
    days = ["MON","TUE","WED","THU","FRI","SAT","SUN"]
    cw = 1.66; x0, y0 = MX, 1.9
    for i,d in enumerate(days):
        x = x0 + i*(cw+0.04)
        hd = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y0), Inches(cw), Inches(0.45))
        hd.fill.solid(); hd.fill.fore_color.rgb=C(INK); hd.line.fill.background(); _noshadow(hd)
        text(s, x, y0+0.06, cw, 0.34, d, font=MONO, size=11, color=CREAM, bold=True, align="c")
        cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y0+0.45), Inches(cw), Inches(3.4))
        cell.fill.solid(); cell.fill.fore_color.rgb=C(WHITE if i%2==0 else NEUT); cell.line.color.rgb=C(INK); cell.line.width=Pt(1); _noshadow(cell)
        text(s, x+0.1, y0+0.55, cw-0.2, 3.2, "[ pillar ]\n[ format ]\n\n[ hook ]", font=BODY, size=10.5, color=MUT, leading=1.2)
    text(s, MX, 5.65, 11.9, 0.7, "Rule of thumb: 3-5 posts/week. Protect 2 slots for reels (reach), fill the rest with story + carousel (bond).",
         font=BODY, size=14, color=INK, leading=1.2)
    footer(s, "SOCIAL")

def s_social_formats():
    s = slide(); head(s, "03 · SOCIAL  ·  MEDIA LIBRARY", "Post format library", accent=YELLOW)
    media(s, MX, 1.8, 2.6, 4.5, "REEL", "9:16")
    text(s, MX, 6.35, 2.6, 0.3, "1080×1920 · reach", font=MONO, size=10, color=MUT, align="c")
    media(s, 3.55, 1.8, 3.5, 4.4, "CAROUSEL", "4:5")
    text(s, 3.55, 6.3, 3.5, 0.3, "1080×1350 · save + teach", font=MONO, size=10, color=MUT, align="c")
    media(s, 7.35, 1.8, 2.6, 4.5, "STORY", "9:16")
    text(s, 7.35, 6.35, 2.6, 0.3, "1080×1920 · daily bond", font=MONO, size=10, color=MUT, align="c")
    card(s, 10.2, 1.8, 2.4, 4.5, fill=YELLOW, line=INK, shadow=INK)
    text(s, 10.4, 2.0, 2.0, 0.4, "SPEC", font=MONO, size=11, color=INK, bold=True, spacing=1)
    bullets(s, 10.4, 2.45, 2.05, 3.6, ["Safe margins 10-14%","Caption ≤ 2 lines on-frame","Logo mark bottom-left","End reels on the logo outro","Trending audio in-app"], size=12, gap=9, mfill=INK, color=INK)
    footer(s, "SOCIAL")

def s_social_voice():
    s = slide(); head(s, "03 · SOCIAL  ·  DEEP-DIVE", "Voice & caption formula", accent=YELLOW)
    card(s, MX, 1.8, 6.0, 4.5, fill=CARD)
    text(s, MX+0.28, 2.0, 5.5, 0.4, "THE CAPTION RECIPE", font=MONO, size=11, color=MUT, bold=True, spacing=1)
    steps = ["Line 1 = the search phrase a fan would type","Deliver ONE useful thing, no throat-clearing",
             "Plain talk, no 'the algorithm', no hype words","Close with 'Save this' on teaching posts",
             "4-6 tight hashtags, no wall of tags","No 'link in bio' on educational posts"]
    bullets(s, MX+0.28, 2.5, 5.4, 3.6, steps, size=13.5, gap=10, mfill=PINK)
    card(s, 7.0, 1.8, 5.6, 4.5, fill=INK, line=INK, shadow=None)
    text(s, 7.25, 2.0, 5.1, 0.4, "VOICE IN ONE LINE", font=MONO, size=11, color=YELLOW, bold=True, spacing=1)
    text(s, 7.25, 2.45, 5.1, 1.4, "“The friend who actually knows the music business and tells you straight.”",
         font=HEAD, size=24, color=CREAM, bold=True, leading=1.08)
    text(s, 7.25, 4.15, 5.1, 0.35, "SWAP THESE WORDS", font=MONO, size=10.5, color=YELLOW, bold=True, spacing=1)
    text(s, 7.25, 4.55, 5.1, 1.6, "'algorithm' → what listeners do\n'blow up' → reach the right people\n'content' → the work / the post",
         font=MONO, size=13, color=CREAM, leading=1.35)
    footer(s, "SOCIAL")

def s_social_calendar():
    s = slide(); head(s, "03 · SOCIAL  ·  WORKSHEET", "30-day content calendar", accent=YELLOW)
    cols = ["DATE","PILLAR","FORMAT","HOOK / FIRST LINE","CTA","STATUS"]
    widths = [1.3, 1.7, 1.6, 4.4, 1.6, 1.3]; x0, y0 = MX, 1.85
    x = x0
    for c,w in zip(cols,widths):
        h = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y0), Inches(w), Inches(0.5))
        h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.25); _noshadow(h)
        text(s, x+0.1, y0+0.09, w-0.15, 0.34, c, font=MONO, size=10, color=CREAM, bold=True, spacing=0.3)
        x += w
    for r in range(6):
        y=y0+0.5+r*0.66; x=x0; fill=WHITE if r%2==0 else NEUT
        for w in widths:
            cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.66))
            cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(0.75); _noshadow(cell)
            x += w
    text(s, MX, y0+0.5+6*0.66+0.12, 11.9, 0.4, "Duplicate rows as needed. Batch a month in one sitting, film in one day.", font=BODY, size=12.5, color=MUT)
    footer(s, "SOCIAL")

# =============================================================================
# SECTION 4 — RELEASE
# =============================================================================
def s_rel_timeline():
    s = slide(); head(s, "04 · RELEASE  ·  CORE", "8-week rollout", accent=MINT)
    weeks = [("WK -8",VIOLET,"Master + art locked. Distributor upload."),
             ("WK -6",PINK,"Pre-save live. Pitch to editorial (needs 4 wk lead)."),
             ("WK -4",YELLOW,"Teasers begin. Snippet + canvas + press shots out."),
             ("WK -2",MINT,"Reel countdown. DM the list. Curator pitches sent."),
             ("WK 0",PINK,"Drop day. Every channel, same day. Thank the fans."),
             ("WK +2",VIOLET,"Follow-through: live take, reactions, playlist adds.")]
    rowh = 0.72; y0 = 1.9
    for i,(wk,col,body) in enumerate(weeks):
        y = y0 + i*rowh
        tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MX), Inches(y), Inches(1.2), Inches(0.56))
        tag.fill.solid(); tag.fill.fore_color.rgb=C(col); tag.line.color.rgb=C(INK); tag.line.width=Pt(2); _round(tag,0.12); _noshadow(tag)
        text(s, MX, y+0.13, 1.2, 0.34, wk, font=MONO, size=12, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        bar_w = 2.0 + i*1.55
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MX+1.4), Inches(y+0.08), Inches(bar_w), Inches(0.42))
        bar.fill.solid(); bar.fill.fore_color.rgb=C(NEUT); bar.line.color.rgb=C(INK); bar.line.width=Pt(1.25); _round(bar,0.3); _noshadow(bar)
        text(s, MX+1.6, y+0.11, min(bar_w+3.0, 10.0), 0.4, body, font=BODY, size=12.5, color=INK, anchor="m")
    footer(s, "RELEASE")

def s_rel_checklist():
    s = slide(); head(s, "04 · RELEASE  ·  WORKSHEET", "Pre-release checklist", accent=MINT)
    groups = [("MUSIC",MINT,["Final master (loudness-checked)","Instrumental + clean versions","ISRC + UPC assigned"]),
              ("VISUAL",VIOLET,["Cover art 3000×3000","Spotify Canvas (9:16 loop)","3 press shots + 5 snippet clips"]),
              ("COPY",PINK,["Bio (3 lengths)","Pitch note for curators","Metadata + credits + splits"]),
              ("DISTRIBUTION",YELLOW,["Uploaded 4 wks ahead","Pre-save link live","Editorial pitch submitted"])]
    cw, ch, gx, gy = 5.85, 2.05, 0.3, 0.28; x0, y0 = MX, 1.8
    for i,(name,col,rows) in enumerate(groups):
        r,cx = divmod(i,2); x=x0+cx*(cw+gx); y=y0+r*(ch+gy)
        card(s, x, y, cw, ch, fill=CARD)
        chip(s, x+0.24, y+0.2, name, fill=col, size=10)
        bullets(s, x+0.24, y+0.68, cw-0.5, 1.3, rows, marker="▢", size=12.5, gap=6, mfill=INK)
        x2 = x+3.15
        footer_dummy = None
    footer(s, "RELEASE")

def s_rel_assets():
    s = slide(); head(s, "04 · RELEASE  ·  MEDIA", "Assets per release", accent=MINT)
    media(s, MX, 1.8, 3.0, 3.0, "COVER ART", "1:1 · 3000px")
    media(s, 3.9, 1.8, 2.5, 4.5, "CANVAS", "9:16 loop")
    media(s, 6.55, 1.8, 3.0, 2.1, "PRESS SHOT", "3:2")
    media(s, 6.55, 4.1, 3.0, 2.2, "SNIPPET CLIPS ×5", "9:16")
    card(s, 9.75, 1.8, 2.85, 4.5, fill=MINT, line=INK, shadow=INK)
    text(s, 9.95, 2.0, 2.5, 0.4, "ALSO SHIP", font=MONO, size=11, color=INK, bold=True, spacing=1)
    bullets(s, 9.95, 2.45, 2.5, 3.6, ["Pre-save link","Lyric / visualiser","3 hook variations","EPK one-pager","Fan-DM message"], size=12.5, gap=11, mfill=INK, color=INK)
    footer(s, "RELEASE")

def s_rel_strategy():
    s = slide(); head(s, "04 · RELEASE  ·  DECISION", "Single, waterfall, or EP?", accent=MINT)
    opt = [("SINGLE",PINK,"One track, full push.","Best when: building the habit, testing a sound, limited budget."),
           ("WATERFALL",VIOLET,"Singles that re-collect into an EP.","Best when: you have 3-4 strong tracks and want compounding pre-saves."),
           ("EP / PROJECT",MINT,"Drop as a body of work.","Best when: there's a narrative, press angle, or a live show to anchor it.")]
    cw, gx = 3.86, 0.28; x = MX
    for name,col,tag,body in opt:
        card(s, x, 1.85, cw, 4.4, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.85), Inches(cw), Inches(0.75))
        top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, 1.99, cw, 0.5, name, font=HEAD, size=22, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x+0.28, 2.85, cw-0.56, 0.8, tag, font=HEAD, size=16, color=INK, bold=True, leading=1.05)
        text(s, x+0.28, 3.75, cw-0.56, 2.2, body, font=BODY, size=13.5, color=MUT, leading=1.2)
        x += cw+gx
    footer(s, "RELEASE")

def s_rel_budget():
    s = slide(); head(s, "04 · RELEASE  ·  WORKSHEET", "Release budget", accent=MINT)
    text(s, MX, 1.55, 11.9, 0.4, "Sample split of a ₹30,000 single budget. Move the numbers, don't add zeros before there's an audience.", font=BODY, size=13.5, color=MUT)
    rows = [("Mix & master",MINT,"₹10,000","33%"),("Cover art + canvas",VIOLET,"₹5,000","17%"),
            ("Content shoot day",YELLOW,"₹6,000","20%"),("Paid amplification",PINK,"₹5,000","17%"),
            ("Curator / PR credits",VIOLET,"₹4,000","13%")]
    y = 2.15
    for name,col,amt,pct in rows:
        card(s, MX, y, 8.5, 0.72, fill=CARD, shadow=INK, off=0.05)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(MX+0.25), Inches(y+0.24), Inches(0.24), Inches(0.24))
        dot.fill.solid(); dot.fill.fore_color.rgb=C(col); dot.line.color.rgb=C(INK); dot.line.width=Pt(1.5); _noshadow(dot)
        text(s, MX+0.68, y+0.16, 4.5, 0.4, name, font=BODY, size=15, color=INK, anchor="m")
        text(s, MX+5.3, y+0.16, 1.6, 0.4, amt, font=MONO, size=15, color=INK, bold=True, anchor="m")
        barw = 1.9*(int(pct[:-1])/33)
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MX+6.9), Inches(y+0.22), Inches(max(0.3,barw)), Inches(0.28))
        bar.fill.solid(); bar.fill.fore_color.rgb=C(col); bar.line.color.rgb=C(INK); bar.line.width=Pt(1); _round(bar,0.4); _noshadow(bar)
        y += 0.82
    card(s, 9.4, 2.15, 3.2, 3.55, fill=INK, shadow=None, line=INK)
    text(s, 9.62, 2.4, 2.8, 0.4, "TOTAL", font=MONO, size=11, color=YELLOW, bold=True, spacing=1)
    text(s, 9.62, 2.8, 2.8, 0.9, "₹30,000", font=HEAD, size=38, color=CREAM, bold=True)
    text(s, 9.62, 3.8, 2.8, 1.8, "Per single. Scale down for the first, up once a track proves it can travel.",
         font=BODY, size=13, color="D9CFE6", leading=1.2)
    footer(s, "RELEASE")

# =============================================================================
# SECTION 5 — PROMOTION
# =============================================================================
def s_promo_map():
    s = slide(); head(s, "05 · PROMOTION  ·  CORE", "Earned, owned, paid", accent=PINK)
    cols = [("OWNED",MINT,"You control it.",["Your socials + list","Spotify profile","Website / EPK","WhatsApp / Discord"]),
            ("EARNED",VIOLET,"You have to win it.",["Playlist adds","Press + blogs","Curator features","Fan word-of-mouth"]),
            ("PAID",PINK,"You buy reach.",["Meta / IG ads","Creator seeding","Playlist pitch credits","Boosted reels"])]
    cw, gx = 3.86, 0.28; x = MX
    for name,col,sub,rows in cols:
        card(s, x, 1.8, cw, 4.5, fill=CARD)
        top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(cw), Inches(0.95))
        top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, 1.95, cw, 0.45, name, font=HEAD, size=22, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x, 2.4, cw, 0.35, sub, font=BODY, size=12.5, color=INK if col in (YELLOW,MINT) else WHITE, align="c", italic=True)
        bullets(s, x+0.35, 3.0, cw-0.6, 3.0, rows, size=14, gap=13, mfill=col)
        x += cw+gx
    footer(s, "PROMOTION")

def s_promo_pitch():
    s = slide(); head(s, "05 · PROMOTION  ·  DEEP-DIVE", "PR & playlist pitching", accent=PINK)
    card(s, MX, 1.8, 6.3, 4.5, fill=CARD)
    text(s, MX+0.28, 2.0, 5.8, 0.4, "THE PITCH, IN 5 LINES", font=MONO, size=11, color=MUT, bold=True, spacing=1)
    bullets(s, MX+0.28, 2.5, 5.75, 3.6, ["Who you are, in one line they can quote","The track + release date + a private link",
        "Why it fits THEIR list / beat (be specific)","One proof point: a number or a name","A clear, small ask. Make yes easy."], size=14, gap=12, mfill=PINK)
    card(s, 7.15, 1.8, 5.45, 2.15, fill=MINT, line=INK, shadow=INK)
    text(s, 7.4, 2.0, 5.0, 0.35, "DO", font=MONO, size=11, color=INK, bold=True, spacing=1)
    bullets(s, 7.4, 2.4, 4.9, 1.4, ["Pitch 4 weeks early","Personalise the first line","Follow up once, politely"], marker="✓", size=12.5, gap=5, mfill=INK, color=INK)
    card(s, 7.15, 4.15, 5.45, 2.15, fill=YELLOW, line=INK, shadow=INK)
    text(s, 7.4, 4.35, 5.0, 0.35, "DON'T", font=MONO, size=11, color=INK, bold=True, spacing=1)
    bullets(s, 7.4, 4.75, 4.9, 1.4, ["Mass-BCC a template","Attach a 30MB file","Chase daily or guilt-trip"], marker="✗", size=12.5, gap=5, mfill=PINK, color=INK)
    footer(s, "PROMOTION")

def s_promo_curator():
    s = slide(); head(s, "05 · PROMOTION  ·  IIM ADVANTAGE", "Curator & reviewer outreach", accent=PINK)
    card(s, MX, 1.75, 5.2, 2.3, fill=VIOLET, line=INK, shadow=INK)
    text(s, MX+0.3, 2.0, 4.7, 0.4, "THE IIM REVIEWER DB", font=MONO, size=11, color=YELLOW, bold=True, spacing=1)
    text(s, MX+0.3, 2.45, 4.7, 0.9, "900+ curators", font=HEAD, size=40, color=WHITE, bold=True)
    text(s, MX+0.3, 3.35, 4.7, 0.6, "Editorial, podcasts, playlists, YouTube: filter by free/paid, South-Asian, cost.",
         font=BODY, size=13, color="EDE7F6", leading=1.15)
    bullets(s, 6.7, 1.85, 5.9, 2.2, ["Shortlist 15-20 that genuinely fit the lane",
        "Prioritise the 17 South-Asian curators first","Track free vs paid (SubmitHub credit ≈ ₹85)",
        "Personalise line 1, the DB gives you the angle"], size=14, gap=11, mfill=VIOLET)
    # mini tracker
    cols = ["CURATOR","TYPE","FREE/PAID","SENT","REPLY"]; widths=[3.4,2.1,2.1,2.0,2.3]; x0,y0=MX,4.35
    x=x0
    for c,w in zip(cols,widths):
        h=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y0),Inches(w),Inches(0.45)); h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.25); _noshadow(h)
        text(s, x+0.12,y0+0.07,w-0.15,0.34,c,font=MONO,size=10,color=CREAM,bold=True,spacing=0.3); x+=w
    for r in range(2):
        y=y0+0.45+r*0.6; x=x0; fill=WHITE if r%2==0 else NEUT
        for w in widths:
            cell=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(0.6)); cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(0.75); _noshadow(cell); x+=w
    footer(s, "PROMOTION")

def s_promo_paid():
    s = slide(); head(s, "05 · PROMOTION  ·  DEEP-DIVE", "Paid amplification", accent=PINK)
    card(s, MX, 1.8, 6.0, 4.5, fill=CARD)
    text(s, MX+0.28, 2.0, 5.5, 0.4, "THE ONLY 3 CAMPAIGNS WORTH RUNNING", font=MONO, size=10.5, color=MUT, bold=True, spacing=0.5)
    items = [("Reel boost","Put ₹ behind the reel that already over-performed organically. Don't boost a dud."),
             ("Pre-save / profile","Drive cheap traffic to the follow, not the stream. Build the base."),
             ("Retarget engagers","Warm audience who watched 50%+. Cheapest converters you have.")]
    y=2.5
    for t,d in items:
        chip(s, MX+0.28, y, t.upper(), fill=PINK, size=10)
        text(s, MX+0.28, y+0.42, 5.4, 0.7, d, font=BODY, size=13, color=INK, leading=1.15); y+=1.25
    card(s, 7.0, 1.8, 5.6, 4.5, fill=INK, shadow=None, line=INK)
    text(s, 7.25, 2.0, 5.1, 0.4, "GUARDRAILS", font=MONO, size=11, color=YELLOW, bold=True, spacing=1)
    bullets(s, 7.25, 2.5, 5.1, 3.6, ["Never buy streams or fake followers","Start at ₹200/day, kill losers fast",
        "Judge on cost-per-follow, not views","Creative wins, not budget: test 3 hooks","Paid amplifies a good post, it can't save a bad one"], size=14, gap=13, mfill=YELLOW, color=CREAM)
    footer(s, "PROMOTION")

def s_promo_kpis():
    s = slide(); head(s, "05 · PROMOTION  ·  DASHBOARD", "KPIs & tracking", accent=PINK)
    text(s, MX, 1.55, 11.9, 0.4, "Fill monthly. Watch the trend, not any single number.", font=BODY, size=14, color=MUT)
    cards = [("MONTHLY LISTENERS",PINK,"[ 0 ]","▲ vs last"),("FOLLOWERS",VIOLET,"[ 0 ]","across all"),
             ("SAVES / STREAM",MINT,"[ 0% ]","fan intent"),("PLAYLIST ADDS",YELLOW,"[ 0 ]","this release"),
             ("EMAIL / DM LIST",PINK,"[ 0 ]","owned reach"),("TICKETS / MERCH",VIOLET,"[ ₹0 ]","real money")]
    cw, ch, gx, gy = 3.86, 1.9, 0.28, 0.3; x0, y0 = MX, 2.05
    for i,(k,col,v,note) in enumerate(cards):
        r,cx = divmod(i,3); x=x0+cx*(cw+gx); y=y0+r*(ch+gy)
        card(s, x, y, cw, ch, fill=CARD)
        bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(0.14))
        bar.fill.solid(); bar.fill.fore_color.rgb=C(col); bar.line.fill.background(); _round(bar,0.4); _noshadow(bar)
        text(s, x+0.28, y+0.28, cw-0.5, 0.35, k, font=MONO, size=10.5, color=MUT, bold=True, spacing=0.3)
        text(s, x+0.28, y+0.65, cw-0.5, 0.9, v, font=HEAD, size=36, color=INK, bold=True)
        text(s, x+0.28, y+1.5, cw-0.5, 0.35, note, font=BODY, size=12, color=col, italic=True)
        x += 0
    footer(s, "PROMOTION")

# =============================================================================
# SECTION 6 — DISTRIBUTION
# =============================================================================
def s_dist_compare():
    s = slide(); head(s, "06 · DISTRIBUTION  ·  DECISION", "Distributor comparison", accent=VIOLET)
    cols = ["DISTRIBUTOR","COST MODEL","KEEP ROYALTIES","BEST FOR"]
    widths = [3.0, 3.2, 2.4, 3.3]; x0, y0 = MX, 1.85
    data = [("DistroKid","Flat annual","100%","Volume: unlimited releases"),
            ("TuneCore","Per-release/yr","100%","Fewer drops, want support"),
            ("CD Baby","One-time / release","91%","Pay once, publishing add-on"),
            ("Believe / Amuse","Free + revenue share","varies","Growth deals, advances"),
            ("[ your pick ]","[ … ]","[ … ]","[ … ]")]
    x=x0
    for c,w in zip(cols,widths):
        h=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y0),Inches(w),Inches(0.5)); h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.5); _noshadow(h)
        text(s, x+0.15,y0+0.08,w-0.2,0.35,c,font=MONO,size=11,color=CREAM,bold=True,spacing=0.4); x+=w
    for r,row in enumerate(data):
        y=y0+0.5+r*0.75; x=x0; fill=WHITE if r%2==0 else NEUT
        for w,val in zip(widths,row):
            cell=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(0.75)); cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(1); _noshadow(cell)
            text(s, x+0.15,y+0.2,w-0.2,0.4,val,font=BODY,size=12.5,color=INK); x+=w
    footer(s, "DISTRIBUTION")

def s_dist_checklist():
    s = slide(); head(s, "06 · DISTRIBUTION  ·  WORKSHEET", "DSP delivery checklist", accent=VIOLET)
    groups=[("METADATA",VIOLET,["Exact artist name (matches profiles)","Correct release date + timezone","Genre + language tags","ISRC per track, UPC per release"]),
            ("RIGHTS",PINK,["Songwriter splits agreed in writing","Producer / feature credits listed","Samples cleared","Publishing / PRO registered"]),
            ("PROFILES",MINT,["Spotify for Artists + Apple claimed","Canvas + artist pick queued","Bio + photo synced everywhere","Pre-save / smartlink live"])]
    cw, gx = 3.86, 0.28; x = MX
    for name,col,rows in groups:
        card(s, x, 1.8, cw, 4.5, fill=CARD)
        chip(s, x+0.24, 2.0, name, fill=col, size=10)
        bullets(s, x+0.24, 2.55, cw-0.5, 3.6, rows, marker="▢", size=13, gap=13, mfill=INK)
        x += cw+gx
    footer(s, "DISTRIBUTION")

def s_dist_money():
    s = slide(); head(s, "06 · DISTRIBUTION  ·  DEEP-DIVE", "Where the money comes from", accent=VIOLET)
    text(s, MX, 1.55, 11.9, 0.4, "Streaming is discovery. Real income is stacked from five streams, ranked by what actually pays in India.", font=BODY, size=14, color=MUT)
    rows=[("LIVE SHOWS",PINK,"The engine. Tickets, festival fees, private gigs.","₹₹₹"),
          ("SYNC & LICENSING",VIOLET,"Ads, film, OTT, games. One placement > a million streams.","₹₹₹"),
          ("MERCH",YELLOW,"High margin, sold hardest at shows and drops.","₹₹"),
          ("DIRECT FANS",MINT,"Bandcamp, memberships, tips, superfan bundles.","₹₹"),
          ("STREAMING",PINK,"₹30-50k per million plays. Discovery, not rent.","₹")]
    y=2.15
    for name,col,body,weight in rows:
        card(s, MX, y, 11.9, 0.8, fill=CARD, shadow=INK, off=0.05)
        bar=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MX), Inches(y), Inches(0.16), Inches(0.8)); bar.fill.solid(); bar.fill.fore_color.rgb=C(col); bar.line.fill.background(); _round(bar,0.4); _noshadow(bar)
        text(s, MX+0.4, y+0.15, 3.2, 0.5, name, font=HEAD, size=17, color=INK, bold=True, anchor="m")
        text(s, MX+3.8, y+0.15, 6.8, 0.5, body, font=BODY, size=13.5, color=MUT, anchor="m")
        text(s, MX+11.0, y+0.15, 0.9, 0.5, weight, font=HEAD, size=20, color=col, bold=True, anchor="m", align="r")
        y += 0.9
    footer(s, "DISTRIBUTION")

def s_dist_splits():
    s = slide(); head(s, "06 · DISTRIBUTION  ·  WORKSHEET", "Sync-readiness & splits", accent=VIOLET)
    card(s, MX, 1.8, 6.0, 4.5, fill=CARD)
    text(s, MX+0.28, 2.0, 5.5, 0.4, "SYNC-READY?", font=MONO, size=11, color=MUT, bold=True, spacing=1)
    bullets(s, MX+0.28, 2.5, 5.5, 3.6, ["Clean instrumental + stems exist","All samples cleared, splits in writing",
        "You control (or can grant) master + publishing","Registered with a PRO / publisher","One-line sync pitch + 30-sec edit ready"], marker="▢", size=13.5, gap=13, mfill=VIOLET)
    # splits table
    cols=["WRITER / ROLE","SHARE %","SIGNED?"]; widths=[3.4,1.7,1.5]; x0,y0=7.15,1.9
    x=x0
    for c,w in zip(cols,widths):
        h=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y0),Inches(w),Inches(0.5)); h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.5); _noshadow(h)
        text(s, x+0.12,y0+0.08,w-0.15,0.34,c,font=MONO,size=10.5,color=CREAM,bold=True,spacing=0.3); x+=w
    for r in range(5):
        y=y0+0.5+r*0.62; x=x0; fill=WHITE if r%2==0 else NEUT
        for w in widths:
            cell=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(0.62)); cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(0.75); _noshadow(cell); x+=w
    text(s, 7.15, y0+0.5+5*0.62+0.15, 5.3, 0.8, "Splits must total 100% and be signed before release. This one worksheet prevents most royalty disputes.",
         font=BODY, size=12.5, color=MUT, leading=1.2)
    footer(s, "DISTRIBUTION")

# =============================================================================
# SECTION 7 — COMPONENT LIBRARY
# =============================================================================
def s_lib_1():
    s = slide(); head(s, "COMPONENT LIBRARY  ·  1 / 3", "Blocks: copy onto any slide", accent=MINT)
    # big stat
    card(s, MX, 1.75, 3.7, 2.2, fill=CARD)
    bar=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MX),Inches(1.75),Inches(3.7),Inches(0.14)); bar.fill.solid(); bar.fill.fore_color.rgb=C(PINK); bar.line.fill.background(); _round(bar,0.4); _noshadow(bar)
    text(s, MX+0.28, 2.0, 3.2, 0.35, "STAT BLOCK", font=MONO, size=10, color=MUT, bold=True)
    text(s, MX+0.28, 2.35, 3.2, 0.9, "1M+", font=HEAD, size=44, color=INK, bold=True)
    text(s, MX+0.28, 3.3, 3.2, 0.5, "streams to earn ₹40k. Once.", font=BODY, size=13, color=MUT)
    # quote card
    card(s, 4.75, 1.75, 3.7, 2.2, fill=YELLOW, line=INK, shadow=INK)
    text(s, 5.0, 1.95, 3.2, 0.3, "QUOTE CARD", font=MONO, size=10, color=INK, bold=True)
    text(s, 5.0, 2.3, 3.2, 1.5, "“The line worth putting on a poster.”", font=HEAD, size=22, color=INK, bold=True, leading=1.05, anchor="m")
    # callout
    card(s, 8.9, 1.75, 3.7, 2.2, fill=INK, shadow=None, line=INK)
    text(s, 9.15, 1.95, 3.2, 0.3, "TIP / CALLOUT", font=MONO, size=10, color=YELLOW, bold=True)
    text(s, 9.15, 2.35, 3.2, 1.4, "A short, sharp note the reader shouldn't miss. One idea only.", font=BODY, size=14, color=CREAM, leading=1.2)
    # checklist
    card(s, MX, 4.15, 5.9, 2.15, fill=CARD)
    text(s, MX+0.28, 4.35, 5.4, 0.3, "CHECKLIST BLOCK", font=MONO, size=10, color=MUT, bold=True)
    bullets(s, MX+0.28, 4.75, 5.4, 1.4, ["First thing to tick","Second thing to tick","Third thing to tick"], marker="▢", size=13.5, gap=8, mfill=INK)
    # chips row
    card(s, 6.7, 4.15, 5.9, 2.15, fill=CARD)
    text(s, 6.95, 4.35, 5.4, 0.3, "CHIPS / PILLS", font=MONO, size=10, color=MUT, bold=True)
    x=6.95
    for lbl,col in [("PINK",PINK),("YELLOW",YELLOW),("MINT",MINT),("VIOLET",VIOLET)]:
        w=chip(s, x, 4.8, lbl, fill=col, size=11); x+=w+0.15
    x=6.95
    for lbl in ["LABEL","STATUS","TAG"]:
        w=chip(s, x, 5.4, lbl, fill=WHITE, size=11); x+=w+0.15
    footer(s, "LIBRARY")

def s_lib_2():
    s = slide(); head(s, "COMPONENT LIBRARY  ·  2 / 3", "Layouts", accent=MINT)
    # 2-col
    text(s, MX, 1.7, 3.0, 0.3, "TWO COLUMN", font=MONO, size=10, color=MUT, bold=True)
    card(s, MX, 2.05, 2.85, 1.9, fill=CARD); card(s, MX+3.0, 2.05, 2.85, 1.9, fill=CARD)
    # 3-col
    text(s, 7.0, 1.7, 3.0, 0.3, "THREE COLUMN", font=MONO, size=10, color=MUT, bold=True)
    for i in range(3): card(s, 7.0+i*1.95, 2.05, 1.75, 1.9, fill=CARD)
    # persona card
    text(s, MX, 4.2, 3.0, 0.3, "PERSONA CARD", font=MONO, size=10, color=MUT, bold=True)
    card(s, MX, 4.5, 3.7, 1.9, fill=CARD)
    media(s, MX+0.2, 4.7, 1.1, 1.1, "AVA", "1:1")
    text(s, MX+1.5, 4.72, 2.0, 0.4, "[ Name ]", font=HEAD, size=16, color=INK, bold=True)
    bullets(s, MX+1.5, 5.15, 2.0, 1.1, ["[ trait ]","[ trait ]"], size=11.5, gap=4, mfill=VIOLET)
    # timeline row
    text(s, 4.75, 4.2, 4.0, 0.3, "TIMELINE ROW", font=MONO, size=10, color=MUT, bold=True)
    for i,(wk,col) in enumerate([("-4",PINK),("-2",VIOLET),("0",MINT),("+2",YELLOW)]):
        tag=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.75+i*1.95),Inches(4.55),Inches(0.85),Inches(0.5)); tag.fill.solid(); tag.fill.fore_color.rgb=C(col); tag.line.color.rgb=C(INK); tag.line.width=Pt(2); _round(tag,0.12); _noshadow(tag)
        text(s, 4.75+i*1.95,4.63,0.85,0.34,"WK "+wk,font=MONO,size=10,color=INK if col in (YELLOW,MINT) else WHITE,bold=True,align="c")
        bar=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.65+i*1.95),Inches(4.63),Inches(1.05),Inches(0.34)); bar.fill.solid(); bar.fill.fore_color.rgb=C(NEUT); bar.line.color.rgb=C(INK); bar.line.width=Pt(1); _round(bar,0.4); _noshadow(bar)
    footer(s, "LIBRARY")

def s_lib_3():
    s = slide(); head(s, "COMPONENT LIBRARY  ·  3 / 3", "Matrix, table, media, CTA", accent=MINT)
    # 2x2
    text(s, MX, 1.7, 3.0, 0.3, "2×2 MATRIX", font=MONO, size=10, color=MUT, bold=True)
    card(s, MX, 2.05, 3.0, 3.0, fill=WHITE, shadow=None, radius=0.02)
    for (x1,y1,x2,y2) in [(MX,3.55,MX+3.0,3.55),(MX+1.5,2.05,MX+1.5,5.05)]:
        ln=s.shapes.add_connector(2, Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=C(INK); ln.line.width=Pt(1.5); _noshadow(ln)
    # table
    text(s, 4.4, 1.7, 3.0, 0.3, "TABLE", font=MONO, size=10, color=MUT, bold=True)
    cols=["COL","COL","COL"]; widths=[1.4,1.4,1.4]; x0,y0=4.4,2.05
    x=x0
    for c,w in zip(cols,widths):
        h=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y0),Inches(w),Inches(0.45)); h.fill.solid(); h.fill.fore_color.rgb=C(INK); h.line.color.rgb=C(INK); h.line.width=Pt(1.25); _noshadow(h)
        text(s, x+0.12,y0+0.06,w-0.15,0.34,c,font=MONO,size=10,color=CREAM,bold=True); x+=w
    for r in range(4):
        y=y0+0.45+r*0.55; x=x0; fill=WHITE if r%2==0 else NEUT
        for w in widths:
            cell=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(0.55)); cell.fill.solid(); cell.fill.fore_color.rgb=C(fill); cell.line.color.rgb=C(INK); cell.line.width=Pt(0.75); _noshadow(cell); x+=w
    # media placeholders
    text(s, 8.9, 1.7, 3.7, 0.3, "MEDIA FRAMES", font=MONO, size=10, color=MUT, bold=True)
    media(s, 8.9, 2.05, 1.7, 1.4, "16:9", "")
    media(s, 10.75, 2.05, 0.85, 1.4, "9:16", "")
    media(s, 11.75, 2.05, 0.85, 0.85, "1:1", "")
    # CTA
    text(s, 8.9, 3.7, 3.7, 0.3, "CTA / NEXT STEPS", font=MONO, size=10, color=MUT, bold=True)
    card(s, 8.9, 4.05, 3.7, 1.0, fill=PINK, line=INK, shadow=INK)
    text(s, 9.15, 4.2, 3.2, 0.7, "[ The one action → ]", font=HEAD, size=18, color=WHITE, bold=True, anchor="m")
    # do/dont pair
    text(s, MX, 5.35, 5.0, 0.3, "DO / DON'T PAIR", font=MONO, size=10, color=MUT, bold=True)
    card(s, MX, 5.65, 2.4, 0.85, fill=MINT, line=INK, shadow=None)
    text(s, MX+0.2, 5.78, 2.0, 0.6, "✓ do this", font=BODY, size=13, color=INK, bold=True, anchor="m")
    card(s, MX+2.6, 5.65, 2.4, 0.85, fill=YELLOW, line=INK, shadow=None)
    text(s, MX+2.8, 5.78, 2.0, 0.6, "✗ not this", font=BODY, size=13, color=INK, bold=True, anchor="m")
    footer(s, "LIBRARY")

# =============================================================================
# SECTION 8 — CLOSE
# =============================================================================
def s_next():
    s = slide(); head(s, "WRAP", "Next steps", accent=PINK)
    steps=[("THIS WEEK",PINK,"[ the 1-2 moves that unblock everything else ]"),
           ("THIS MONTH",VIOLET,"[ the release or campaign in motion ]"),
           ("THIS QUARTER",MINT,"[ the bigger bet: tour, EP, sync push ]")]
    cw, gx = 3.86, 0.28; x = MX
    for k,col,v in steps:
        card(s, x, 1.9, cw, 3.9, fill=CARD)
        top=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(1.9),Inches(cw),Inches(0.7)); top.fill.solid(); top.fill.fore_color.rgb=C(col); top.line.color.rgb=C(INK); top.line.width=Pt(2); _round(top,0.07); _noshadow(top)
        text(s, x, 2.02, cw, 0.5, k, font=HEAD, size=18, color=INK if col in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x+0.28, 2.9, cw-0.56, 2.7, v, font=BODY, size=15, color=INK, leading=1.25)
        x += cw+gx
    footer(s, "CLOSE")

def s_signoff():
    s = slide(bg=INK, bar=PINK)
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(MX), Inches(2.2), height=Inches(1.4))
    text(s, MX+1.8, 2.35, 9.0, 1.4, "Indie Music India", font=HEAD, size=52, color=CREAM, bold=True)
    text(s, MX, 4.4, 11.0, 0.6, "The playbook, the tools, and the people who use them.", font=BODY, size=20, color="D9CFE6")
    x = MX
    for lbl in ["indiemusicindia.com","@indiemusicindia.co","hello@indiemusicindia.com"]:
        w = chip(s, x, 5.4, lbl, fill=YELLOW, size=10); x += w + 0.18
    return s

# =============================================================================
# EDITORIAL LAYERS — thesis, moodboard, references, takeaways
# =============================================================================
def s_statement(kicker, line, accent):
    """A single-sentence thesis slide (dark, full-bleed)."""
    s = slide(bg=INK, bar=accent)
    chip(s, MX, 1.35, kicker, fill=accent)
    text(s, MX, 2.0, 11.9, 3.4, line, font=HEAD, size=44, color=CREAM, bold=True, leading=1.1, anchor="m")
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def s_moodboard(kicker, title, accent, tiles, note):
    s = slide(); head(s, kicker, title, accent)
    cw, ch, gx, gy = 2.75, 1.75, 0.24, 0.24; x0, y0 = MX, 1.8
    for i,(lbl,ratio) in enumerate(tiles):
        r,cx = divmod(i,3); x=x0+cx*(cw+gx); y=y0+r*(ch+gy)
        media(s, x, y, cw, ch, lbl, ratio)
    card(s, 9.55, 1.8, 3.05, 3.98, fill=accent, line=INK, shadow=INK)
    text(s, 9.78, 2.02, 2.6, 0.35, "WHAT THIS SAYS", font=MONO, size=10.5, color=INK, bold=True, spacing=0.5)
    text(s, 9.78, 2.5, 2.6, 3.1, note, font=BODY, size=13.5, color=INK, leading=1.22)
    footer(s, kicker.split("·")[0].strip())
    return s

def s_references(kicker, title, accent, refs):
    s = slide(); head(s, kicker, title, accent)
    text(s, MX, 1.55, 11.9, 0.4, "The north stars. Not to copy, to triangulate: take one thing from each, become none of them.", font=BODY, size=14, color=MUT)
    cw, gx = 3.86, 0.28; x = MX
    cols = [PINK, VIOLET, MINT]
    for i,(name, borrow) in enumerate(refs):
        card(s, x, 2.1, cw, 4.2, fill=CARD)
        media(s, x+0.28, 2.35, cw-0.56, 1.7, "REFERENCE", "3:2")
        text(s, x+0.28, 4.15, cw-0.56, 0.5, name, font=HEAD, size=18, color=INK, bold=True)
        text(s, x+0.28, 4.65, cw-0.56, 0.3, "WHAT WE BORROW", font=MONO, size=9.5, color=cols[i], bold=True, spacing=0.5)
        text(s, x+0.28, 5.0, cw-0.56, 1.2, borrow, font=BODY, size=13, color=INK, leading=1.18)
        x += cw+gx
    footer(s, kicker.split("·")[0].strip())
    return s

def s_takeaways(section, accent, points):
    s = slide(); head(s, section+"  ·  RECAP", "Takeaways", accent)
    cw, gx = 3.86, 0.28; x = MX
    for i,(t,d) in enumerate(points):
        card(s, x, 1.95, cw, 4.1, fill=CARD)
        num = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+0.28), Inches(2.2), Inches(0.7), Inches(0.7))
        num.fill.solid(); num.fill.fore_color.rgb=C(accent); num.line.color.rgb=C(INK); num.line.width=Pt(2); _round(num,0.2); _noshadow(num)
        text(s, x+0.28, 2.3, 0.7, 0.5, str(i+1), font=HEAD, size=26, color=INK if accent in (YELLOW,MINT) else WHITE, bold=True, align="c")
        text(s, x+0.28, 3.15, cw-0.56, 0.9, t, font=HEAD, size=19, color=INK, bold=True, leading=1.05)
        text(s, x+0.28, 4.15, cw-0.56, 1.8, d, font=BODY, size=13.5, color=MUT, leading=1.2)
        x += cw+gx
    footer(s, section)
    return s

# =============================================================================
# BLANKS & SECTIONS — empty branded canvases to build from
# =============================================================================
def s_section_template(accent):
    """A generic, retitle-able section divider. Duplicate and relabel per client."""
    s = slide(bg=INK, bar=accent)
    text(s, MX, 1.5, 4.5, 2.0, "00", font=HEAD, size=150, color=accent, bold=True, wrap=False)
    text(s, MX, 3.55, 11.0, 1.4, "[ SECTION TITLE ]", font=HEAD, size=68, color=CREAM, bold=True)
    text(s, MX, 4.85, 9.0, 0.6, "[ one-line purpose of this section ]", font=BODY, size=20, color="D9CFE6")
    x = MX
    for _ in range(3):
        w = chip(s, x, 5.7, "[ WHAT'S INSIDE ]", fill=accent, fg=INK, size=12); x += w + 0.18
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def s_blank_title():
    s = slide(); chip(s, MX, 0.55, "[ LABEL ]", fill=PINK)
    text(s, MX, 0.97, 11.9, 0.9, "[ Slide title ]", font=HEAD, size=34, color=INK, bold=True)
    card(s, MX, 1.8, 11.9, 4.5, fill=CARD)
    text(s, MX+0.35, 2.05, 11.2, 0.4, "[ Body: type here, or paste a block from the Component Library ]", font=BODY, size=15, color=MUT)
    footer(s, "BLANK CANVAS")
    return s

def s_blank_twocol():
    s = slide(); chip(s, MX, 0.55, "[ LABEL ]", fill=VIOLET)
    text(s, MX, 0.97, 11.9, 0.9, "[ Two-column title ]", font=HEAD, size=34, color=INK, bold=True)
    card(s, MX, 1.8, 5.85, 4.5, fill=CARD)
    text(s, MX+0.3, 2.05, 5.3, 0.4, "[ Left column ]", font=BODY, size=15, color=MUT)
    card(s, 6.75, 1.8, 5.85, 4.5, fill=CARD)
    text(s, 7.05, 2.05, 5.3, 0.4, "[ Right column ]", font=BODY, size=15, color=MUT)
    footer(s, "BLANK CANVAS")
    return s

def s_blank_media():
    s = slide(); chip(s, MX, 0.55, "[ LABEL ]", fill=MINT)
    text(s, MX, 0.97, 11.9, 0.9, "[ Media-led title ]", font=HEAD, size=34, color=INK, bold=True)
    media(s, MX, 1.8, 7.4, 4.5, "DROP IMAGE / VIDEO", "16:9")
    card(s, 8.3, 1.8, 4.3, 4.5, fill=CARD)
    text(s, 8.6, 2.05, 3.7, 0.4, "[ Caption / notes ]", font=BODY, size=15, color=MUT)
    footer(s, "BLANK CANVAS")
    return s

def s_blank_plain():
    s = slide()
    text(s, MX, 0.5, 6.0, 0.4, "[ blank canvas, number stays for reference ]", font=MONO, size=11, color=MUT)
    footer(s, "BLANK CANVAS")
    return s

def _blank_head(s, title, kfill=YELLOW):
    chip(s, MX, 0.55, "[ LABEL ]", fill=kfill)
    text(s, MX, 0.97, 11.9, 0.9, title, font=HEAD, size=34, color=INK, bold=True)

def s_blank_titlecover():
    s = slide(bg=INK, bar=PINK)
    chip(s, MX, 1.2, "[ KICKER ]", fill=YELLOW)
    text(s, MX, 1.85, 11.9, 2.3, "[ Big cover title ]", font=HEAD, size=60, color=CREAM, bold=True, leading=1.0)
    text(s, MX, 4.5, 11.0, 0.8, "[ Subtitle or one-line promise ]", font=BODY, size=20, color="D9CFE6")
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def s_blank_threecol():
    s = slide(); _blank_head(s, "[ Three-column title ]")
    cw=3.86; gx=0.28; x=MX
    for _ in range(3):
        card(s, x, 1.8, cw, 4.5, fill=CARD)
        text(s, x+0.28, 2.05, cw-0.56, 0.4, "[ Column ]", font=BODY, size=15, color=MUT)
        x+=cw+gx
    footer(s, "BLANK CANVAS"); return s

def s_blank_statement():
    s = slide(bg=INK, bar=VIOLET)
    chip(s, MX, 1.35, "[ KICKER ]", fill=VIOLET)
    text(s, MX, 2.0, 11.9, 3.4, "[ One-sentence statement goes here. ]", font=HEAD, size=44, color=CREAM, bold=True, leading=1.1, anchor="m")
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def s_blank_quote():
    s = slide(); _blank_head(s, "[ Quote / testimonial title ]", kfill=PINK)
    card(s, MX, 1.9, 11.9, 4.4, fill=YELLOW, line=INK, shadow=INK)
    text(s, MX+0.5, 2.2, 11.0, 3.0, "“[ The quote goes here. ]”", font=HEAD, size=34, color=INK, bold=True, leading=1.08, anchor="m")
    text(s, MX+0.5, 5.55, 11.0, 0.5, "[ Name, role ]", font=MONO, size=13, color=INK)
    footer(s, "BLANK CANVAS"); return s

def s_blank_stats():
    s = slide(); _blank_head(s, "[ Metrics / numbers title ]", kfill=MINT)
    cw=2.78; gx=0.13; x=MX
    for col in [PINK, VIOLET, MINT, YELLOW]:
        card(s, x, 1.9, cw, 2.0, fill=CARD); _topbar(s, x, 1.9, cw, col)
        text(s, x+0.24, 2.2, cw-0.4, 0.9, "[ 0 ]", font=HEAD, size=40, color=INK, bold=True)
        text(s, x+0.24, 3.25, cw-0.4, 0.4, "[ LABEL ]", font=MONO, size=10.5, color=MUT, bold=True)
        x+=cw+gx
    footer(s, "BLANK CANVAS"); return s

def s_blank_compare():
    s = slide(); _blank_head(s, "[ Comparison / vs title ]", kfill=VIOLET)
    card(s, MX, 1.9, 5.5, 4.4, fill=CARD)
    text(s, MX+0.3, 2.15, 5.0, 0.4, "[ Option A ]", font=HEAD, size=20, color=INK, bold=True)
    card(s, 7.1, 1.9, 5.5, 4.4, fill=CARD)
    text(s, 7.4, 2.15, 5.0, 0.4, "[ Option B ]", font=HEAD, size=20, color=INK, bold=True)
    vs = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.36), Inches(3.7), Inches(0.9), Inches(0.9))
    vs.fill.solid(); vs.fill.fore_color.rgb=C(PINK); vs.line.color.rgb=C(INK); vs.line.width=Pt(2); _noshadow(vs)
    text(s, 6.36, 3.88, 0.9, 0.5, "VS", font=HEAD, size=20, color=WHITE, bold=True, align="c")
    footer(s, "BLANK CANVAS"); return s

def s_blank_gallery():
    s = slide(); _blank_head(s, "[ Gallery / moodboard title ]", kfill=PINK)
    tw=3.86; th=2.15; gx=0.28; gy=0.2; x0,y0=MX,1.85
    for i in range(6):
        r,c=divmod(i,3); media(s, x0+c*(tw+gx), y0+r*(th+gy), tw, th, "IMAGE", "")
    footer(s, "BLANK CANVAS"); return s

def s_blank_fullbleed():
    s = slide()
    media(s, 0.3, 0.3, SW-0.6, SH-0.6, "DROP FULL-BLEED IMAGE / VIDEO: SEND TO BACK, ADD TEXT OVER", "fill the frame")
    chip(s, MX, SH-1.15, "[ CAPTION OVER IMAGE ]", fill=YELLOW)
    return s

# =============================================================================
# BRAND PARTNERSHIP PITCH — a second deliverable type (templatised from a real
# title-partner deck): pitch a project/artist to brands for sponsorship.
# =============================================================================
def _leftbar(s, x, y, h, col):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.14), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb=C(col); b.line.fill.background(); _round(b,0.5); _noshadow(b)

def _topbar(s, x, y, w, col, h=0.14):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb=C(col); b.line.fill.background(); _round(b,0.4); _noshadow(b)

def s_bp_cover():
    s = slide(bg=INK, bar=YELLOW)
    chip(s, MX, 0.85, "PARTNER PROPOSAL  ·  CATEGORY EXCLUSIVE", fill=YELLOW)
    text(s, MX, 1.5, 11.9, 1.5, "[ Project / film name ]", font=HEAD, size=64, color=CREAM, bold=True)
    text(s, MX, 2.95, 11.0, 0.5, "[ Theme  ·  Theme  ·  Theme ]", font=BODY, size=22, color="D9CFE6", italic=True)
    card(s, MX, 3.95, 8.4, 1.25, fill=INK, line=YELLOW, lw=1.75, shadow=None)
    text(s, MX+0.3, 4.13, 7.8, 0.35, "PARTNER PROPOSAL FOR", font=MONO, size=11, color=YELLOW, bold=True, spacing=1)
    text(s, MX+0.3, 4.5, 7.8, 0.6, "[ Brand ]", font=HEAD, size=30, color=CREAM, bold=True)
    text(s, MX, 5.5, 11.9, 0.5, "Featuring [ talent · talent ]      Director [ name ]      DOP [ name ]", font=BODY, size=14, color="9C93A8")
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

def s_bp_credibility():
    s = slide(); head(s, "BRAND PITCH  ·  CREDIBILITY", "The act, by the numbers", accent=YELLOW)
    stats=[("[ 0 ]","SPOTIFY MONTHLY",PINK),("[ 0 ]","YOUTUBE VIEWS",VIOLET),("[ 0 ]","INSTAGRAM",MINT),("[ 0 ]","FESTIVALS",YELLOW)]
    cw=2.78; gx=0.13; x=MX
    for v,k,col in stats:
        card(s,x,1.75,cw,1.55,fill=CARD); _topbar(s,x,1.75,cw,col)
        text(s,x+0.24,1.98,cw-0.4,0.75,v,font=HEAD,size=32,color=INK,bold=True)
        text(s,x+0.24,2.78,cw-0.4,0.4,k,font=MONO,size=10,color=MUT,bold=True,spacing=0.3)
        x+=cw+gx
    q=[("“[ Pull quote that positions the act ]”","·  [ Publication, year ]"),("“[ Second credibility quote ]”","·  [ Publication ]")]
    cw2=5.85; x=MX
    for quote,src in q:
        card(s,x,3.6,cw2,2.7,fill=NEUT,line=INK,shadow=INK)
        text(s,x+0.3,3.85,cw2-0.6,1.7,quote,font=HEAD,size=20,color=INK,bold=True,leading=1.1)
        text(s,x+0.3,5.72,cw2-0.6,0.4,src,font=MONO,size=11,color=MUT)
        x+=cw2+0.2
    footer(s,"BRAND PITCH")
    return s

def s_bp_concept():
    s=slide(); head(s,"BRAND PITCH  ·  CREATIVE","The concept & visual language",accent=YELLOW)
    card(s,MX,1.75,5.6,4.55,fill=INK,shadow=None,line=INK)
    text(s,MX+0.3,1.98,5.0,0.35,"THE CONCEPT",font=MONO,size=11,color=YELLOW,bold=True,spacing=1)
    text(s,MX+0.3,2.4,5.0,1.8,"“[ The idea in one line a brand can repeat. ]”",font=HEAD,size=25,color=CREAM,bold=True,leading=1.08)
    text(s,MX+0.3,4.3,5.0,0.35,"THE MOTIF",font=MONO,size=11,color=YELLOW,bold=True,spacing=1)
    text(s,MX+0.3,4.7,5.0,1.4,"[ The single visual thread that carries it from first frame to last. ]",font=BODY,size=14,color="D9CFE6",leading=1.2)
    tiles=[("SET / WORLD","4:3"),("LIGHT / GRADE","4:3"),("TEXTURE / OBJECT","4:3"),("MOTIF DETAIL","4:3")]
    tw=2.7; th=2.15; x0=6.6; y0=1.75
    for i,(lbl,r) in enumerate(tiles):
        rr,cc=divmod(i,2); media(s,x0+cc*(tw+0.14),y0+rr*(th+0.11),tw,th,lbl,r)
    footer(s,"BRAND PITCH")
    return s

def s_bp_cast():
    s=slide(); head(s,"BRAND PITCH  ·  TALENT","The cast & featured artists",accent=YELLOW)
    cw=3.86; gx=0.28; x=MX; cols=[PINK,VIOLET,MINT]
    for i in range(3):
        card(s,x,1.8,cw,4.5,fill=CARD)
        media(s,x+0.28,2.05,cw-0.56,1.85,"PORTRAIT","3:4")
        text(s,x+0.28,4.0,cw-0.56,0.5,"[ Name ]",font=HEAD,size=19,color=INK,bold=True)
        text(s,x+0.28,4.5,cw-0.56,0.35,"[ Role / known for ]",font=BODY,size=12,color=MUT)
        chip(s,x+0.28,4.95,"[ HEADLINE STAT ]",fill=cols[i],size=9)
        text(s,x+0.28,5.5,cw-0.56,0.7,"[ One credit or award line. ]",font=BODY,size=12,color=INK,leading=1.15)
        x+=cw+gx
    footer(s,"BRAND PITCH")
    return s

def s_bp_assets():
    s=slide(); head(s,"BRAND PITCH  ·  THE OFFER","Assets, not impressions",accent=YELLOW)
    text(s,MX,1.55,11.9,0.4,"What the brand owns outright, forever. No approval rounds, no expiry.",font=BODY,size=14,color=MUT)
    items=[("Full photo library","Cast shoot images to run as Meta / Google / print / OOH, forever.",PINK),
           ("3-5 brand cutdowns","Ad-ready edits to run on YouTube / Meta independently.",VIOLET),
           ("Title card + placement","Presented by [Brand]. Product lit and framed, editorial-grade.",YELLOW),
           ("Social at launch","Cast and artist posts and reels at release.",MINT),
           ("EPK & press","Named as Title Partner in the press kit and release.",PINK),
           ("Billboard creative","City OOH creative, yours to run, no approval.",VIOLET)]
    cw=3.86; ch=1.6; gx=0.28; gy=0.22; x0,y0=MX,2.05
    for i,(t,d,col) in enumerate(items):
        r,cx=divmod(i,3); x=x0+cx*(cw+gx); y=y0+r*(ch+gy)
        card(s,x,y,cw,ch,fill=CARD); _leftbar(s,x,y,ch,col)
        text(s,x+0.32,y+0.18,cw-0.55,0.4,t,font=HEAD,size=16,color=INK,bold=True)
        text(s,x+0.32,y+0.62,cw-0.55,ch-0.72,d,font=BODY,size=11.5,color=MUT,leading=1.14)
    footer(s,"BRAND PITCH")
    return s

def s_bp_audience():
    s=slide(); head(s,"BRAND PITCH  ·  REACH","Audience & the maths",accent=YELLOW)
    segs=[("[ Core geography ]","[ who they are, why they buy ]",PINK),
          ("[ Diaspora ]","[ international, high-intent cultural buyers ]",VIOLET),
          ("[ Cultural youth ]","[ aspirational, culturally conscious ]",MINT)]
    y=1.8
    for t,d,col in segs:
        card(s,MX,y,6.0,1.32,fill=CARD); _leftbar(s,MX,y,1.32,col)
        text(s,MX+0.32,y+0.2,5.5,0.4,t,font=HEAD,size=17,color=INK,bold=True)
        text(s,MX+0.32,y+0.68,5.5,0.5,d,font=BODY,size=12.5,color=MUT)
        y+=1.5
    card(s,7.05,1.8,5.55,1.55,fill=INK,shadow=None,line=INK)
    text(s,7.3,1.98,5.0,0.3,"COMBINED REACH, YEAR 1",font=MONO,size=10.5,color=YELLOW,bold=True,spacing=0.5)
    text(s,7.3,2.3,3.4,0.8,"[ 0M ]",font=HEAD,size=40,color=CREAM,bold=True)
    text(s,9.9,2.55,2.6,0.7,"views\nCPV [ ₹0 ]",font=BODY,size=13,color="D9CFE6",leading=1.15)
    comp=[("vs festival sponsor","[ ₹ ]  ·  ends Sunday, own nothing",PINK),
          ("vs celebrity influencer","[ ₹ ]  ·  stops when the campaign stops",VIOLET),
          ("This partnership","owned forever, compounds",MINT)]
    y=3.55
    for k,v,col in comp:
        card(s,7.05,y,5.55,0.82,fill=CARD,shadow=INK,off=0.05); _leftbar(s,7.05,y,0.82,col)
        text(s,7.35,y+0.13,5.0,0.35,k,font=HEAD,size=14,color=INK,bold=True)
        text(s,7.35,y+0.46,5.0,0.3,v,font=MONO,size=10.5,color=MUT)
        y+=0.94
    footer(s,"BRAND PITCH")
    return s

def s_bp_deliverables():
    s=slide(); head(s,"BRAND PITCH  ·  DELIVERABLES","Content architecture",accent=YELLOW)
    cols=[("PRIMARY FILM",PINK,["Main film (master)","Director's cut","Trailer 90-120s","Teaser 45-60s"]),
          ("SHORT-FORM",VIOLET,["4x 30s chapter reels","4x 15s platform reels","7-chapter still series","Social card set"]),
          ("EDITORIAL & ARTIST",YELLOW,["Character posters","Ensemble poster","Product / costume editorial","BTS photo series"]),
          ("PRESS & BRAND",MINT,["BTS documentary","Electronic press kit","Brand cutdowns","Press release"])]
    cw=2.8; gx=0.12; x=MX
    for name,col,rows in cols:
        card(s,x,1.8,cw,4.5,fill=CARD); _topbar(s,x,1.8,cw,col,h=0.62)
        text(s,x+0.2,1.92,cw-0.4,0.4,name,font=HEAD,size=14,color=INK if col in (YELLOW,MINT) else WHITE,bold=True,anchor="m")
        bullets(s,x+0.24,2.65,cw-0.45,3.4,rows,size=12.5,gap=12,mfill=col)
        x+=cw+gx
    footer(s,"BRAND PITCH")
    return s

def s_bp_tiers():
    s=slide(); head(s,"BRAND PITCH  ·  PACKAGES","Three ways to partner",accent=YELLOW)
    text(s,MX,1.55,11.9,0.4,"Category exclusive. Choose the tier that fits the brand's ambition and budget.",font=BODY,size=14,color=MUT)
    tiers=[("TITLE PARTNER","[ ₹ ]",PINK,["The film exists because of you","Title card + song-title option","Full photo library + 3-5 cutdowns","All film and social assets","Category exclusive, no approvals"]),
           ("ASSOCIATE","[ ₹ ]",VIOLET,["Alongside a title partner","Secondary logo on trailer + reels","Selected images + 2 cutdowns","Cast story posts at launch","EPK credit as associate"]),
           ("INTEGRATION","[ ₹ ]",MINT,["Product inside the film's world","Choose an integration zone","Dedicated editorial stills","Cast post if worn by talent","EPK product credit"])]
    cw=3.86; gx=0.28; x=MX
    for name,price,col,rows in tiers:
        card(s,x,2.1,cw,4.2,fill=CARD); _topbar(s,x,2.1,cw,col,h=0.9)
        text(s,x+0.28,2.22,cw-0.5,0.4,name,font=HEAD,size=17,color=INK if col in (YELLOW,MINT) else WHITE,bold=True)
        text(s,x+0.28,2.6,cw-0.5,0.35,price,font=MONO,size=15,color=INK if col in (YELLOW,MINT) else WHITE,bold=True)
        bullets(s,x+0.28,3.2,cw-0.55,2.9,rows,size=12.5,gap=10,mfill=col)
        x+=cw+gx
    footer(s,"BRAND PITCH")
    return s

def s_bp_next():
    s=slide(); head(s,"BRAND PITCH  ·  CLOSE","What you provide + next step",accent=YELLOW)
    card(s,MX,1.8,6.0,4.5,fill=CARD)
    text(s,MX+0.3,2.0,5.5,0.35,"WHAT YOU PROVIDE (ALL TIERS)",font=MONO,size=11,color=MUT,bold=True,spacing=0.5)
    bullets(s,MX+0.3,2.5,5.4,3.5,["High-res logo (PNG + SVG)","Exact name spelling","Brand colour hex codes","One-line descriptor (20 words max)","Product or props for the shoot"],size=14,gap=14,mfill=PINK)
    card(s,7.05,1.8,5.55,4.5,fill=YELLOW,line=INK,shadow=INK)
    text(s,7.35,2.0,5.0,0.35,"NEXT STEP",font=MONO,size=11,color=INK,bold=True,spacing=1)
    text(s,7.35,2.5,5.0,1.6,"A 30-minute call with [ names / roles ] to confirm the tier, integration scope, and proceed to contract.",font=HEAD,size=21,color=INK,bold=True,leading=1.1)
    text(s,7.35,4.6,5.0,0.35,"THE ONE NON-NEGOTIABLE",font=MONO,size=10.5,color=INK,bold=True,spacing=0.5)
    text(s,7.35,5.0,5.0,1.1,"No approval rounds on creative, across every tier. Stated before the contract is signed.",font=BODY,size=13.5,color=INK,leading=1.18)
    footer(s,"BRAND PITCH")
    return s

def s_bp_proof():
    s=slide(); head(s,"BRAND PITCH  ·  WHY NOW","Market context & proof",accent=YELLOW)
    cards=[("THE FESTIVAL ENDS",PINK,"[ Festival sponsorships vanish when the gates close. A film lives on digital forever. ]","REFERENCE · [ festival, source, date ]"),
           ("THE WAVE YOU MISSED",VIOLET,"[ The cultural moment the brand was adjacent to but never inside. Name the proof: titles, streams, awards. ]","REFERENCE · [ titles, year ]"),
           ("PLATFORMS ARE BETTING",MINT,"[ Government, YouTube, Spotify are backing independent creative. This project is that bet. ]","REFERENCE · [ MoU / payout, date ]")]
    cw=3.86; gx=0.28; x=MX
    for name,col,body,ref in cards:
        card(s,x,1.85,cw,4.4,fill=CARD); _topbar(s,x,1.85,cw,col)
        text(s,x+0.28,2.15,cw-0.55,0.5,name,font=HEAD,size=17,color=INK,bold=True)
        text(s,x+0.28,2.75,cw-0.55,2.4,body,font=BODY,size=14,color=INK,leading=1.22)
        text(s,x+0.28,5.55,cw-0.55,0.6,ref,font=MONO,size=9,color=MUT,leading=1.15)
        x+=cw+gx
    footer(s,"BRAND PITCH")
    return s

def s_bp_close():
    s=slide(bg=INK,bar=YELLOW)
    text(s, MX, 2.1, 11.9, 2.0, "[ Your product. Their faces. Yours forever. ]", font=HEAD, size=44, color=CREAM, bold=True, leading=1.1, anchor="m")
    text(s, MX, 4.15, 11.9, 0.9, "[ One line on what this really is: a cultural property built to produce campaign-grade imagery a brand owns and runs forever. ]", font=BODY, size=16, color="D9CFE6", leading=1.2)
    x=MX
    for lbl in ["[ contact name · role ]","[ phone ]","[ email · site · @handle ]"]:
        w=chip(s,x,5.5,lbl,fill=YELLOW,size=10); x+=w+0.18
    if MARK.exists(): s.shapes.add_picture(str(MARK), Inches(SW-1.1), Inches(0.42), height=Inches(0.5))
    return s

# =============================================================================
def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    # Section 0 — ORIENTATION
    s_cover(); s_howto(); s_brand(); s_pipeline(); s_projectcover()
    # BLANKS & SECTIONS (up front, for quick grab)
    divider("00","BLANKS & SECTIONS",VIOLET,"Empty branded canvases. Duplicate, retitle, build. Grab these first.",
        ["Sections","Covers","Layouts","Statement","Quote","Stats","Compare","Gallery"])
    s_section_template(PINK); s_section_template(MINT); s_section_template(VIOLET)
    s_blank_titlecover(); s_blank_title(); s_blank_twocol(); s_blank_threecol()
    s_blank_media(); s_blank_fullbleed(); s_blank_gallery()
    s_blank_statement(); s_blank_quote(); s_blank_stats(); s_blank_compare(); s_blank_plain()
    # Section 1 — IDENTITY
    divider(*SECTIONS[0], ["Thesis","Snapshot","Origin","Moodboard","References","Audit","Takeaways"])
    s_statement("01 · IDENTITY", "Before anyone can love the music, they have to be able to describe the artist in one sentence.", PINK)
    s_identity_snapshot(); s_identity_origin(); s_identity_sonic()
    s_moodboard("01 · IDENTITY  ·  MOODBOARD", "Artist moodboard",
        PINK, [("TEXTURE / GRAIN","1:1"),("COLOUR STORY","1:1"),("TYPE / LOGO","1:1"),("SPACE / VENUE","4:3"),("STYLING / WARDROBE","3:4"),("MOTION / GRADE","16:9")],
        "Pull 6-8 images that feel like the record before you brief a shoot. If two tiles fight, the identity isn't settled yet.")
    s_identity_visual()
    s_references("01 · IDENTITY  ·  REFERENCES", "Creative north stars", PINK,
        [("[ Reference artist ]","[ e.g. their honesty in a vocal, the way they use silence ]"),
         ("[ Reference artist ]","[ e.g. their visual restraint, one colour done well ]"),
         ("[ Reference / non-music ]","[ e.g. a film's grade, a designer's grid ]")])
    s_identity_audit()
    s_takeaways("IDENTITY", PINK, [
        ("One sentence, first","If the artist can't be said in a line, nothing downstream lands cleanly."),
        ("Specific beats flattering","A place, a turn, a sound. Never a pile of adjectives."),
        ("Define the 'not'","The lane gets sharp from what the artist refuses to be.")])
    # Section 2 — POSITIONING
    divider(*SECTIONS[1], ["Thesis","Statement","Personas","Perceptual map","Wedge","Teardown","Takeaways"])
    s_statement("02 · POSITIONING", "You don't win by being better. You win by being the only one who does a specific thing.", VIOLET)
    s_pos_statement(); s_pos_personas(); s_pos_map(); s_pos_wedge(); s_pos_teardown()
    s_takeaways("POSITIONING", VIOLET, [
        ("Find the empty corner","Aim for the quadrant that's uncrowded and still true to the artist."),
        ("The wedge is un-copyable","If a peer could claim it by next week, it isn't a wedge."),
        ("Write for one persona","Everyone is nobody. Pick the fan you're actually for.")])
    # Section 3 — SOCIAL
    divider(*SECTIONS[2], ["Thesis","Channels","Pillars","Cadence","Formats","Voice","Calendar","Takeaways"])
    s_statement("03 · SOCIAL", "Social isn't promotion. It's the relationship you keep building between releases.", YELLOW)
    s_social_channels(); s_social_pillars(); s_social_cadence(); s_social_formats(); s_social_voice(); s_social_calendar()
    s_takeaways("SOCIAL", YELLOW, [
        ("Every post has a pillar","If it doesn't fit music / person / value / proof, it doesn't go out."),
        ("Protect the reel slots","Reach lives there. Teach one thing and saves carry it further."),
        ("Voice over volume","Sound like a straight-talking friend, not a press release.")])
    # Section 4 — RELEASE
    divider(*SECTIONS[3], ["Thesis","8-wk rollout","Moodboard","Checklist","Assets","Strategy","Budget","Takeaways"])
    s_statement("04 · RELEASE", "A release is a campaign with a date, not a file with a date.", MINT)
    s_rel_timeline()
    s_moodboard("04 · RELEASE  ·  ART DIRECTION", "Campaign moodboard",
        MINT, [("COVER DIRECTION","1:1"),("CANVAS / MOTION","9:16"),("KEY VISUAL","16:9"),("SNIPPET LOOK","9:16"),("PRESS SHOT TONE","3:2"),("TYPE TREATMENT","1:1")],
        "One look across cover, canvas, reels, and press. If the artwork and the reels feel like different acts, the rollout leaks.")
    s_rel_checklist(); s_rel_assets(); s_rel_strategy(); s_rel_budget()
    s_takeaways("RELEASE", MINT, [
        ("Work backwards 8 weeks","Start from the date and reverse-engineer every task to it."),
        ("Editorial needs 4 weeks","Miss the lead time and you're pitching a track that's already out."),
        ("Assets before ads","A great organic post beats a boosted weak one, every time.")])
    # Section 5 — PROMOTION
    divider(*SECTIONS[4], ["Thesis","Earned/owned/paid","Pitching","Curator DB","Paid","KPIs","Takeaways"])
    s_statement("05 · PROMOTION", "No one shares music they had to work to understand. Make the ask easy.", PINK)
    s_promo_map(); s_promo_pitch(); s_promo_curator(); s_promo_paid(); s_promo_kpis()
    s_takeaways("PROMOTION", PINK, [
        ("Own the list","You can be de-ranked from a feed. You can't be de-ranked from an inbox."),
        ("Personalise line one","The first sentence decides whether the rest gets read."),
        ("Judge paid on follows","Cost-per-follow and saves, never raw views.")])
    # Section 6 — DISTRIBUTION
    divider(*SECTIONS[5], ["Thesis","Distributors","DSP checklist","Money flow","Sync & splits","Takeaways"])
    s_statement("06 · DISTRIBUTION", "Streaming is how you get found. It is not how you get paid.", VIOLET)
    s_dist_compare(); s_dist_checklist(); s_dist_money(); s_dist_splits()
    s_takeaways("DISTRIBUTION", VIOLET, [
        ("Splits in writing, pre-release","One signed worksheet prevents most royalty disputes."),
        ("Keep stems clean","Sync is the biggest single cheque. Stay ready to say yes fast."),
        ("Stack five streams","Live, sync, merch, direct, streaming. Streaming is only one of them.")])
    # Section 7 — BRAND PARTNERSHIP PITCH (a second deliverable type)
    divider("07","BRAND PARTNERSHIP",YELLOW,"Turn a project into a brand asset a partner owns forever.",
        ["Cover","Credibility","The offer","Reach","Deliverables","Tiers","Proof"])
    s_statement("07 · BRAND PARTNERSHIP", "Brands don't want another logo on a stage. They want an asset they own forever.", YELLOW)
    s_bp_cover(); s_bp_credibility(); s_bp_concept(); s_bp_cast(); s_bp_assets(); s_bp_audience()
    s_bp_deliverables(); s_bp_tiers(); s_bp_next(); s_bp_proof(); s_bp_close()
    # Section 8 — COMPONENT LIBRARY
    divider("08","COMPONENT LIBRARY",MINT,"The drag-and-drop kit. Copy any block onto your slide.", ["Blocks","Layouts","Matrix / table / media / CTA"])
    s_lib_1(); s_lib_2(); s_lib_3()
    # CLOSE
    s_next(); s_signoff()
    stamp_numbers()
    prs.save(str(OUT))
    print(f"Deck: {len(prs.slides._sldIdLst)} slides -> {OUT}")

if __name__ == "__main__":
    build()
