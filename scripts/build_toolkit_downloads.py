#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_toolkit_downloads.py — neat, branded, watermarked downloads for the toolkit.

For each template in the toolkit, from its pulled .txt source:
  - SHEET-type (budgets, P&L, trackers, expenses...) -> branded editable .xlsx (openpyxl)
  - DOC-type   (agreements, split sheets, email/one-sheet/metadata...) -> branded
    watermarked .pdf (reportlab), fill-in fields as lines, poppy brand, IIM watermark.
Writes toolkit/<slug>/<slug>.(pdf|xlsx). Brand rule: strips em dashes.
Run: python3 scripts/build_toolkit_downloads.py
"""
import re, json
from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                Spacer, HRFlowable, Image, Table, TableStyle)
from reportlab.lib.styles import ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPO = Path(__file__).resolve().parent.parent
SRC = Path.home()/"Downloads"/"iim_toolkit"
FONTS = REPO/"assets"/"fonts"

INK=HexColor("#241B2E"); CREAM=HexColor("#FFF7EE"); PINK=HexColor("#FF4D8D")
MUT=HexColor("#6B6076"); PINKD=HexColor("#D8005E"); LINE=HexColor("#E7DFD3")

SHEET_KW=("budget","profit","loss","p-l","cashflow","cash-flow","expense","revenue",
          "forecast","tracker","matrix","net-worth","finances","bookkeeping","inventory","cost-")

for n,f in (("Bric","BricolageGrotesque.ttf"),("Inter","Inter.ttf"),("Mono","SpaceMono-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(n,str(FONTS/f)))

def norm(txt):
    return txt.replace(" — ", ", ").replace("—", ", ").replace(" – ", ", ").replace("–","-")

def esc(s): return s.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

def src_for(slug):
    p=SRC/f"{slug}.txt"
    if p.exists(): return p
    # manifest slugs may differ (e.g. annual-profit-amp-loss); fuzzy match
    key=slug.replace("-template","").replace("-sheet","").split("-")[0]
    cand=sorted(SRC.glob(f"*{key}*.txt"))
    return cand[0] if cand else None

def parse(txt, title):
    out=[]; para=[]; ul=[]
    def fp():
        nonlocal para
        if para: out.append(("p"," ".join(para))); para=[]
    def fu():
        nonlocal ul
        if ul: out.append(("ul",ul[:])); ul=[]
    seen=False
    for raw in txt.replace("﻿","").split("\n"):
        ln=raw.rstrip()
        s=ln.strip()
        if not s: fp(); fu(); continue
        if len(s)>3 and set(s)<=set("━—–-=_·.|─ "): fp(); fu(); out.append(("hr",None)); continue
        if "Toolkit |" in s: continue
        if "|" in s and s.count("|")>=2:   # table row
            fp(); fu(); out.append(("row",[c.strip() for c in s.split("|")])); continue
        letters=[c for c in s if c.isalpha()]
        is_head=letters and sum(c.isupper() for c in letters)/len(letters)>0.85 and len(s)<70
        if is_head:
            h=s.strip(":")
            if h.upper()==title.upper():
                if not seen: seen=True; continue
                continue
            fp(); fu(); out.append(("h2",h.title() if h.isupper() else h)); continue
        if s[0] in "•-*▪◦" or s.startswith("☐") or s[:3] in ("[ ]","[]"):
            fp(); ul.append(re.sub(r"^[•\-*▪◦☐\[\]\s]+","",s)); continue
        if s.endswith(":") and len(s)<48: fp(); fu(); out.append(("field",s)); continue
        fu(); para.append(s)
    fp(); fu()
    return out

# ---------------- PDF (doc-type) ----------------
def watermark(cv, doc):
    cv.saveState(); cv.setFont("Bric",54); cv.setFillColor(INK); cv.setFillAlpha(0.04)
    cv.translate(A4[0]/2,A4[1]/2); cv.rotate(45); cv.drawCentredString(0,0,"INDIE MUSIC INDIA"); cv.restoreState()
    cv.saveState(); cv.setFont("Mono",7.5); cv.setFillColor(MUT)
    cv.drawString(20*mm,12*mm,"FREE TEMPLATE  ·  INDIEMUSICINDIA.COM")
    cv.drawRightString(A4[0]-20*mm,12*mm,str(doc.page))
    cv.setStrokeColor(PINK); cv.setLineWidth(1.4); cv.line(20*mm,15*mm,A4[0]-20*mm,15*mm); cv.restoreState()

def build_pdf(slug,title,blocks):
    S=lambda **k: ParagraphStyle(k.pop("n"),**k)
    H1=S(n="H1",fontName="Bric",fontSize=24,leading=28,textColor=INK)
    EYE=S(n="EYE",fontName="Mono",fontSize=9,leading=12,textColor=PINKD,spaceAfter=10)
    H2=S(n="H2",fontName="Bric",fontSize=14,leading=18,textColor=INK,spaceBefore=15,spaceAfter=5)
    BODY=S(n="BODY",fontName="Inter",fontSize=10.5,leading=16,textColor=INK,spaceAfter=7)
    FIELD=S(n="FIELD",fontName="Inter",fontSize=10.5,leading=20,textColor=INK,spaceAfter=3)
    LI=S(n="LI",fontName="Inter",fontSize=10.5,leading=15,textColor=INK,leftIndent=12,spaceAfter=3)
    out=REPO/"toolkit"/slug; out.mkdir(parents=True,exist_ok=True); pdf=out/f"{slug}.pdf"
    doc=BaseDocTemplate(str(pdf),pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,
                        topMargin=22*mm,bottomMargin=22*mm,title=title,author="Indie Music India")
    doc.addPageTemplates([PageTemplate(id="m",frames=[Frame(doc.leftMargin,doc.bottomMargin,doc.width,doc.height)],onPage=watermark)])
    st=[]; mark=REPO/"assets"/"brand"/"mark.png"
    if mark.exists(): im=Image(str(mark),14*mm,14*mm); im.hAlign="LEFT"; st+=[im,Spacer(1,4)]
    st.append(Paragraph("FREE TEMPLATE · INDIE MUSIC INDIA",EYE))
    st.append(Paragraph(esc(title),H1)); st.append(HRFlowable(width="100%",thickness=2,color=PINK,spaceBefore=8,spaceAfter=12))
    rowbuf=[]
    def flush_rows():
        nonlocal rowbuf
        if not rowbuf: return
        w=doc.width; ncol=max(len(r) for r in rowbuf); rowbuf=[r+[""]*(ncol-len(r)) for r in rowbuf]
        t=Table([[Paragraph(esc(c),ParagraphStyle("c",fontName="Inter",fontSize=8.5,leading=11,textColor=INK)) for c in r] for r in rowbuf],colWidths=[w/ncol]*ncol)
        t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,LINE),("BACKGROUND",(0,0),(-1,0),HexColor("#FFE1EC")),
            ("FONTNAME",(0,0),(-1,0),"Mono"),("FONTSIZE",(0,0),(-1,0),7.5),("VALIGN",(0,0),(-1,-1),"TOP"),
            ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5),("LEFTPADDING",(0,0),(-1,-1),6)]))
        st.append(Spacer(1,4)); st.append(t); st.append(Spacer(1,8)); rowbuf=[]
    for i,(typ,val) in enumerate(blocks):
        if typ=="row": rowbuf.append(val); continue
        flush_rows()
        if typ=="h2": st.append(Paragraph(esc(val),H2))
        elif typ=="hr": st.append(HRFlowable(width="100%",thickness=0.6,color=LINE,spaceBefore=6,spaceAfter=6))
        elif typ=="p": st.append(Paragraph(esc(val),BODY))
        elif typ=="field": st.append(Paragraph(f"<b>{esc(val)}</b>&nbsp;"+("_"*58),FIELD))
        elif typ=="ul":
            for x in val: st.append(Paragraph(esc(x),LI,bulletText="•"))
    flush_rows()
    st.append(Spacer(1,16)); st.append(HRFlowable(width="100%",thickness=0.6,color=LINE,spaceAfter=6))
    st.append(Paragraph("Free to use from indiemusicindia.com. Share the page, not the file.",
                        ParagraphStyle("f",fontName="Inter",fontSize=8.5,leading=12,textColor=MUT)))
    doc.build(st); return pdf

# ---------------- XLSX (sheet-type) ----------------
def build_xlsx(slug,title,blocks):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Template"
    inkf=Font(name="Calibri",bold=True,color="FFFFFF"); pinkfill=PatternFill("solid",fgColor="FF4D8D")
    inkfill=PatternFill("solid",fgColor="241B2E"); cream=PatternFill("solid",fgColor="FFF7EE")
    hdr=Font(name="Calibri",bold=True,color="FFFFFF"); bold=Font(name="Calibri",bold=True,color="241B2E")
    muted=Font(name="Calibri",italic=True,color="6B6076",size=9)
    thin=Side(style="thin",color="E7DFD3"); border=Border(thin,thin,thin,thin)
    r=1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    c=ws.cell(r,1,"◉ INDIE MUSIC INDIA  ·  FREE TEMPLATE"); c.font=Font(name="Calibri",bold=True,color="D8005E",size=11); r+=1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    c=ws.cell(r,1,title); c.font=Font(name="Calibri",bold=True,color="241B2E",size=16); r+=2
    AMT=("inr","amount","estimated","actual","total","₹","cost","price","value","revenue","expense")
    for typ,val in blocks:
        if typ=="h2":
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
            cc=ws.cell(r,1,val); cc.font=hdr; cc.fill=inkfill; r+=1
        elif typ=="field":
            ws.cell(r,1,val).font=bold; r+=1
        elif typ=="p":
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
            ws.cell(r,1,val).font=muted; r+=1
        elif typ=="ul":
            for x in val: ws.cell(r,2,"• "+x); r+=1
        elif typ=="row":
            is_hdr=any(any(k in cell.lower() for k in AMT) for cell in val) or r==3
            for j,cell in enumerate(val,1):
                cc=ws.cell(r,j,cell); cc.border=border
                if is_hdr: cc.font=Font(bold=True,color="241B2E"); cc.fill=PatternFill("solid",fgColor="FFE1EC")
            hdr_row=r; r+=1
            if is_hdr:
                amt_cols=[j for j,cell in enumerate(val,1) if any(k in cell.lower() for k in AMT)]
                for _ in range(8):
                    for j in range(1,len(val)+1): ws.cell(r,j).border=border
                    r+=1
                if amt_cols:
                    ws.cell(r,1,"TOTAL").font=bold
                    for j in amt_cols:
                        col=openpyxl.utils.get_column_letter(j)
                        ws.cell(r,j,f"=SUM({col}{hdr_row+1}:{col}{r-1})").font=bold
                    r+=1
            r+=1
    widths=[26,22,16,16,22,16]
    for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.sheet_view.showGridLines=False
    footer=ws.cell(r+1,1,"Free to use from indiemusicindia.com. Share the page, not the file."); footer.font=muted
    out=REPO/"toolkit"/slug; out.mkdir(parents=True,exist_ok=True); xp=out/f"{slug}.xlsx"; wb.save(xp); return xp

def main():
    slugs=sorted(p.parent.name for p in REPO.glob("toolkit/*/index.html")
                 if "docs.google.com/document" in p.read_text(encoding="utf-8",errors="ignore"))
    man={d["slug"]:d for d in json.loads((SRC/"manifest.json").read_text())}
    pdfs=xlsxs=miss=0
    for slug in slugs:
        f=src_for(slug)
        if not f: print("  MISSING SRC",slug); miss+=1; continue
        txt=norm(f.read_text(encoding="utf-8",errors="ignore"))
        title=txt.replace("﻿","").split("\n")[0].strip()
        title=title.title() if title.isupper() else title
        blocks=parse(txt,title)
        if any(k in slug for k in SHEET_KW):
            build_xlsx(slug,title,blocks); xlsxs+=1
        else:
            build_pdf(slug,title,blocks); pdfs+=1
    print(f"built {pdfs} PDFs + {xlsxs} XLSX ({miss} missing source)")

if __name__=="__main__":
    main()
